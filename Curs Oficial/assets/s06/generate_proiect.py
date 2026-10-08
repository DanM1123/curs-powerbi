# -*- coding: utf-8 -*-
"""Temă · Sesiunea 6: 4 seturi de date murdare + pagina temei (Curățare · Model · Raport).

Pentru fiecare set de date se generează:
  sesiuni/s06/tema/<companie>/  -> 5 fișiere .xlsx murdare + tema.html + CITESTE_MAI_INTAI.txt
  assets/s06/tema_curat/<companie>/ -> aceleași tabele, curățate (pentru facilitator)

Fiecare problemă de curățare are o singură rezolvare clară în Power Query. Numerele de control
din pagina de proiect (rânduri, coloane, sume, rezultate KPI) se calculează din datele curate.
Rulare: python generate_proiect.py
"""
from __future__ import annotations

import random
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parent
OUT = ROOT.parent.parent / "sesiuni" / "s06" / "tema"
OUT_CURAT = ROOT / "tema_curat"

CAL_START = date(2024, 1, 1)
CAL_END = date(2025, 12, 31)
LUNI = ["", "Ianuarie", "Februarie", "Martie", "Aprilie", "Mai", "Iunie", "Iulie", "August",
        "Septembrie", "Octombrie", "Noiembrie", "Decembrie"]
ZILE = ["Luni", "Marti", "Miercuri", "Joi", "Vineri", "Sambata", "Duminica"]

REGIUNE = {
    "Cluj-Napoca": "Nord-Vest", "Oradea": "Nord-Vest", "Baia Mare": "Nord-Vest",
    "Timisoara": "Vest", "Arad": "Vest",
    "Brasov": "Centru", "Sibiu": "Centru", "Targu Mures": "Centru",
    "Iasi": "Nord-Est", "Suceava": "Nord-Est",
    "Constanta": "Sud-Est", "Galati": "Sud-Est",
    "Craiova": "Sud-Vest",
    "Ploiesti": "Sud", "Pitesti": "Sud",
    "Bucuresti": "Bucuresti-Ilfov", "Voluntari": "Bucuresti-Ilfov",
}
REGIUNE_MURDAR = {"Nord-Vest": "Nord Vest", "Sud-Est": "Sud Est", "Nord-Est": "Nord Est",
                  "Bucuresti-Ilfov": "Bucuresti Ilfov", "Sud-Vest": "Sud Vest"}
COORD = {
    "Cluj-Napoca": (46.7712, 23.6236), "Oradea": (47.0465, 21.9189), "Baia Mare": (47.6567, 23.5850),
    "Timisoara": (45.7489, 21.2087), "Arad": (46.1866, 21.3123), "Brasov": (45.6427, 25.5887),
    "Sibiu": (45.7983, 24.1256), "Targu Mures": (46.5455, 24.5625), "Iasi": (47.1585, 27.6014),
    "Suceava": (47.6514, 26.2556), "Constanta": (44.1598, 28.6348), "Galati": (45.4353, 28.0080),
    "Craiova": (44.3302, 23.7949), "Ploiesti": (44.9365, 26.0129), "Pitesti": (44.8565, 24.8692),
    "Bucuresti": (44.4268, 26.1025), "Voluntari": (44.4925, 26.1764),
}
PF_NUME = ["Ion Popescu", "Maria Ionescu", "Andrei Pop", "Elena Radu", "Mihai Stan", "Ioana Dumitru",
           "Alexandru Matei", "Cristina Lazar", "Bogdan Nistor", "Raluca Toma", "Vlad Constantin",
           "Diana Luca", "George Ilie", "Laura Sandu", "Radu Enache", "Simona Nita", "Victor Marinescu",
           "Ana Cojocaru", "Catalin Moldovan", "Oana Barbu", "Sorin Petrescu", "Irina Voicu",
           "Florin Dobre", "Mirela Stoica", "Dan Ungureanu", "Adina Rusu"]


@dataclass
class Team:
    folder: str
    color: str
    company: str
    tagline: str
    context: tuple[str, ...]
    seed: int
    # nume tabele (= nume fișiere .xlsx și nume tabele în Power BI)
    fact: str
    t_cli: str
    t_prod: str
    t_loc: str
    # coloane cheie / specifice
    c_id: str              # ID tranzacție
    c_data: str
    c_cli: str             # ID client (în fact și în dimensiune)
    c_prod: str
    c_loc: str
    c_qty: str
    c_pret: str
    c_canal: str
    status_ok: str
    status_ko: str
    canale: tuple[tuple[str, float], ...]
    pct_ko: float
    seasonality: tuple[float, ...]   # 12 ponderi pe luni
    # dimensiuni
    cli_nume: str          # numele coloanei de nume din tabelul de clienți
    cli_data: str
    segmente_pf: tuple[str, ...]
    segmente_pj: tuple[str, ...]
    firme: tuple[str, ...]
    p_nume: str
    p_cod: str
    p_extra: tuple[str, str] | None   # (nume coloană, tip) – coloană specifică industriei
    produse: tuple[tuple, ...]        # (cod, nume, categorie, extra, pret, popularitate, (qmin,qmax))
    l_nume: str
    locatii: tuple[tuple, ...]        # (nume, oras, format, suprafata, pondere)
    titles: dict = field(default_factory=dict)
    report: dict = field(default_factory=dict)


# ---------------------------------------------------------------------------
# Cele 4 seturi de date
# ---------------------------------------------------------------------------
TEAMS: list[Team] = [
    Team(
        folder="cafeloop", color="#DC2626", company="CafeLoop", tagline="rețea de cafenele · HoReCa",
        context=(
            "CafeLoop are cafenele în 9 orașe: stradale, în mall, în aeroport și pe campusuri. Vinde cafea, ceai, patiserie, sandvișuri și cafea boabe la pachet.",
            "Managementul vrea un raport de 2 pagini: cum merge rețeaua și ce, cui și pe ce canal se vinde.",
        ),
        seed=20261101,
        fact="Vanzari", t_cli="Clienti", t_prod="Produse", t_loc="Cafenele",
        c_id="ID_Bon", c_data="Data", c_cli="ID_Client", c_prod="ID_Produs", c_loc="ID_Cafenea",
        c_qty="Cantitate", c_pret="Pret_Unitar", c_canal="Canal",
        status_ok="Platit", status_ko="Anulat",
        canale=(("La casa", 0.52), ("Aplicatie", 0.24), ("Livrare", 0.14), ("Corporate", 0.10)),
        pct_ko=0.06,
        seasonality=(1.25, 1.2, 1.05, 0.95, 0.85, 0.75, 0.7, 0.72, 0.95, 1.1, 1.2, 1.35),
        cli_nume="Nume_Client", cli_data="Data_Inregistrare",
        segmente_pf=("Student", "Profesionist", "Turist"), segmente_pj=("Corporate",),
        firme=("SC Tech Hub SRL", "SC Office Park SA", "SC Event Plus SRL", "SC Media Group SRL",
               "SC Law Partners SRL", "SC Startup Lab SRL", "SC Bank Services SA", "SC Coworking One SRL"),
        p_nume="Nume_Produs", p_cod="Cod_Produs", p_extra=None,
        produse=(
            ("CF-01", "Espresso", "Cafea", None, 9.0, 10, (1, 3)),
            ("CF-02", "Cappuccino", "Cafea", None, 14.0, 12, (1, 3)),
            ("CF-03", "Latte", "Cafea", None, 15.0, 11, (1, 3)),
            ("CF-04", "Flat White", "Cafea", None, 16.0, 7, (1, 2)),
            ("CF-05", "Cold Brew", "Cafea", None, 17.0, 5, (1, 2)),
            ("CE-01", "Ceai verde", "Ceai", None, 11.0, 4, (1, 2)),
            ("CE-02", "Ceai de fructe", "Ceai", None, 11.0, 4, (1, 2)),
            ("CE-03", "Chai Latte", "Ceai", None, 15.0, 4, (1, 2)),
            ("PT-01", "Croissant unt", "Patiserie", None, 9.0, 8, (1, 4)),
            ("PT-02", "Briosa ciocolata", "Patiserie", None, 10.0, 6, (1, 4)),
            ("PT-03", "Cheesecake", "Patiserie", None, 18.0, 4, (1, 2)),
            ("PT-04", "Cookie ovaz", "Patiserie", None, 7.0, 5, (1, 5)),
            ("SV-01", "Sandvis pui", "Sandvisuri", None, 22.0, 5, (1, 3)),
            ("SV-02", "Sandvis vegan", "Sandvisuri", None, 21.0, 3, (1, 3)),
            ("SV-03", "Wrap ton", "Sandvisuri", None, 23.0, 3, (1, 2)),
            ("SV-04", "Bagel somon", "Sandvisuri", None, 27.0, 2, (1, 2)),
            ("BA-01", "Cafea boabe Brazilia 250g", "Boabe si accesorii", None, 45.0, 2, (1, 3)),
            ("BA-02", "Cafea boabe Etiopia 250g", "Boabe si accesorii", None, 52.0, 2, (1, 2)),
            ("BA-03", "Cana termos CafeLoop", "Boabe si accesorii", None, 69.0, 1, (1, 2)),
            ("BA-04", "Filtru V60", "Boabe si accesorii", None, 39.0, 1, (1, 2)),
        ),
        l_nume="Nume_Cafenea",
        locatii=(
            ("CafeLoop Cluj Centru", "Cluj-Napoca", "Stradal", 120, 1.4),
            ("CafeLoop Bucuresti Universitate", "Bucuresti", "Campus", 140, 1.6),
            ("CafeLoop Baneasa Mall", "Bucuresti", "Mall", 90, 1.2),
            ("CafeLoop Aeroport Otopeni", "Voluntari", "Aeroport", 70, 1.1),
            ("CafeLoop Timisoara Unirii", "Timisoara", "Stradal", 110, 1.0),
            ("CafeLoop Iasi Palas", "Iasi", "Mall", 95, 0.9),
            ("CafeLoop Brasov Republicii", "Brasov", "Stradal", 85, 0.8),
            ("CafeLoop Constanta Port", "Constanta", "Stradal", 80, 0.6),
            ("CafeLoop Sibiu Campus", "Sibiu", "Campus", 75, 0.6),
            ("CafeLoop Craiova Drive", "Craiova", "Drive-thru", 60, 0.5),
        ),
        titles={"fact": ("Export bonuri · POS CafeLoop", "Generat automat - nu modificati"),
                "cli": ("Clienti aplicatie CafeLoop", "export CRM decembrie 2025"),
                "loc": ("Lista cafenele CafeLoop", "actualizat de echipa de operatiuni"),
                "cal": "Calendar CafeLoop 2024-2025", "prod": "Meniu CafeLoop"},
    ),
    Team(
        folder="travelgo", color="#2563EB", company="TravelGo", tagline="agenție de turism · pachete de vacanță",
        context=(
            "TravelGo vinde pachete de vacanță (city break, plajă, munte, circuite, destinații exotice) prin agenții fizice, website, telefon și parteneri.",
            "Directorul comercial vrea un raport de 2 pagini: când și unde se vând rezervările și ce destinații și clienți aduc valoare.",
        ),
        seed=20261102,
        fact="Rezervari", t_cli="Clienti", t_prod="Pachete", t_loc="Agentii",
        c_id="ID_Rezervare", c_data="Data_Rezervare", c_cli="ID_Client", c_prod="ID_Pachet", c_loc="ID_Agentie",
        c_qty="Nr_Persoane", c_pret="Pret_Persoana", c_canal="Canal",
        status_ok="Confirmata", status_ko="Anulata",
        canale=(("Agentie", 0.38), ("Website", 0.34), ("Telefon", 0.14), ("Partener", 0.14)),
        pct_ko=0.11,
        seasonality=(1.6, 1.5, 1.3, 1.0, 1.1, 1.15, 0.8, 0.6, 0.65, 0.75, 0.8, 0.75),
        cli_nume="Nume_Client", cli_data="Data_Inregistrare",
        segmente_pf=("Familie", "Cuplu", "Single", "Senior"), segmente_pj=("Corporate",),
        firme=("SC Incentive Travel SRL", "SC Pharma Team SA", "SC Auto Parts SRL", "SC IT Solutions SRL",
               "SC Retail Group SA", "SC Energy Plus SRL"),
        p_nume="Nume_Pachet", p_cod="Cod_Pachet", p_extra=("Tara", "Text"),
        produse=(
            ("CB-01", "City break Roma", "City break", "Italia", 1450.0, 6, (1, 4)),
            ("CB-02", "City break Paris", "City break", "Franta", 1690.0, 5, (1, 4)),
            ("CB-03", "City break Viena", "City break", "Austria", 1190.0, 4, (1, 4)),
            ("CB-04", "City break Barcelona", "City break", "Spania", 1590.0, 5, (1, 4)),
            ("PL-01", "Sejur Antalya all inclusive", "Plaja", "Turcia", 2890.0, 9, (2, 5)),
            ("PL-02", "Sejur Creta", "Plaja", "Grecia", 2690.0, 8, (2, 5)),
            ("PL-03", "Sejur Mamaia", "Plaja", "Romania", 1490.0, 6, (2, 5)),
            ("PL-04", "Sejur Mallorca", "Plaja", "Spania", 3190.0, 5, (2, 4)),
            ("MT-01", "Ski Poiana Brasov", "Munte", "Romania", 1290.0, 4, (1, 4)),
            ("MT-02", "Ski Bansko", "Munte", "Bulgaria", 1890.0, 4, (1, 4)),
            ("MT-03", "Ski Alpi austrieci", "Munte", "Austria", 3490.0, 3, (1, 4)),
            ("MT-04", "Drumetie Transfagarasan", "Munte", "Romania", 790.0, 3, (1, 4)),
            ("CI-01", "Circuit Italia de Nord", "Circuit", "Italia", 3990.0, 3, (1, 3)),
            ("CI-02", "Circuit Andaluzia", "Circuit", "Spania", 4290.0, 2, (1, 3)),
            ("CI-03", "Circuit Grecia antica", "Circuit", "Grecia", 3590.0, 2, (1, 3)),
            ("EX-01", "Dubai 7 nopti", "Exotic", "Emiratele Arabe", 5990.0, 3, (1, 4)),
            ("EX-02", "Maldive 7 nopti", "Exotic", "Maldive", 11990.0, 1, (2, 2)),
            ("EX-03", "Thailanda Phuket", "Exotic", "Thailanda", 7490.0, 2, (1, 3)),
            ("EX-04", "Egipt Hurghada", "Exotic", "Egipt", 3290.0, 4, (2, 5)),
            ("EX-05", "Zanzibar 9 nopti", "Exotic", "Tanzania", 9490.0, 1, (2, 2)),
        ),
        l_nume="Nume_Agentie",
        locatii=(
            ("TravelGo Bucuresti Victoriei", "Bucuresti", "Agentie stradala", 85, 1.5),
            ("TravelGo AFI Cotroceni", "Bucuresti", "Mall", 40, 1.2),
            ("TravelGo Cluj Iulius", "Cluj-Napoca", "Mall", 38, 1.1),
            ("TravelGo Timisoara Centru", "Timisoara", "Agentie stradala", 60, 0.9),
            ("TravelGo Iasi Copou", "Iasi", "Agentie stradala", 55, 0.8),
            ("TravelGo Brasov", "Brasov", "Agentie stradala", 50, 0.7),
            ("TravelGo Constanta", "Constanta", "Agentie stradala", 45, 0.6),
            ("TravelGo Oradea Lotus", "Oradea", "Mall", 35, 0.5),
            ("TravelGo Corporate Desk", "Bucuresti", "Corporate", 30, 0.7),
            ("TravelGo Online", "Bucuresti", "Online", None, 1.4),
        ),
        titles={"fact": ("Export rezervari · sistem TravelGo Booking", "Date confidentiale"),
                "cli": ("Baza clienti TravelGo", "export ianuarie 2026"),
                "loc": ("Agentii TravelGo", "coordonate pentru harta"),
                "cal": "Calendar rezervari 2024-2025", "prod": "Catalog pachete TravelGo"},
    ),
    Team(
        folder="petzone", color="#059669", company="PetZone", tagline="rețea de pet shop-uri",
        context=(
            "PetZone vinde hrană, accesorii, produse de îngrijire, sănătate și jucării pentru animale, în magazine fizice, online și pe bază de abonament lunar.",
            "Managerul de rețea vrea un raport de 2 pagini: cum merg magazinele și pentru ce animale și ce categorii se vinde.",
        ),
        seed=20261103,
        fact="Vanzari", t_cli="Clienti", t_prod="Produse", t_loc="Magazine",
        c_id="ID_Bon", c_data="Data", c_cli="ID_Client", c_prod="ID_Produs", c_loc="ID_Magazin",
        c_qty="Cantitate", c_pret="Pret_Unitar", c_canal="Canal",
        status_ok="Platit", status_ko="Returnat",
        canale=(("Magazin", 0.55), ("Online", 0.22), ("Aplicatie", 0.12), ("Abonament", 0.11)),
        pct_ko=0.05,
        seasonality=(0.9, 0.85, 0.95, 1.0, 1.05, 1.05, 1.0, 1.0, 1.0, 1.05, 1.15, 1.4),
        cli_nume="Nume_Client", cli_data="Data_Inregistrare",
        segmente_pf=("Stapan caine", "Stapan pisica", "Mixt"), segmente_pj=("Crescator", "Clinica veterinara"),
        firme=("SC Vet Care SRL", "SC Canisa Regala SRL", "SC Pet Hotel SRL", "SC Clinica Animed SRL",
               "SC Crescatorie Felina SRL", "SC Dog Training SRL"),
        p_nume="Nume_Produs", p_cod="Cod_Produs", p_extra=("Animal", "Text"),
        produse=(
            ("HR-01", "Hrana uscata caini adulti 12kg", "Hrana", "Caine", 289.0, 9, (1, 2)),
            ("HR-02", "Hrana umeda caini 400g", "Hrana", "Caine", 12.0, 8, (2, 12)),
            ("HR-03", "Hrana uscata pisici 4kg", "Hrana", "Pisica", 139.0, 9, (1, 2)),
            ("HR-04", "Plicuri pisici 85g", "Hrana", "Pisica", 5.0, 9, (4, 24)),
            ("HR-05", "Seminte papagali 1kg", "Hrana", "Pasari", 29.0, 3, (1, 3)),
            ("HR-06", "Fulgi pesti 250ml", "Hrana", "Pesti", 24.0, 2, (1, 3)),
            ("AC-01", "Lesa extensibila", "Accesorii", "Caine", 79.0, 4, (1, 1)),
            ("AC-02", "Cusca transport pisica", "Accesorii", "Pisica", 159.0, 2, (1, 1)),
            ("AC-03", "Acvariu 60L", "Accesorii", "Pesti", 449.0, 1, (1, 1)),
            ("AC-04", "Colivie hamsteri", "Accesorii", "Rozatoare", 189.0, 1, (1, 1)),
            ("IN-01", "Sampon caini", "Ingrijire", "Caine", 39.0, 4, (1, 2)),
            ("IN-02", "Nisip pisici 10L", "Ingrijire", "Pisica", 35.0, 7, (1, 4)),
            ("IN-03", "Perie deparazitare", "Ingrijire", "Pisica", 45.0, 3, (1, 2)),
            ("SA-01", "Pipeta antiparazitara caini", "Sanatate", "Caine", 69.0, 5, (1, 3)),
            ("SA-02", "Vitamine pisici", "Sanatate", "Pisica", 49.0, 3, (1, 2)),
            ("SA-03", "Tratament apa acvariu", "Sanatate", "Pesti", 32.0, 1, (1, 2)),
            ("JC-01", "Minge cauciuc", "Jucarii", "Caine", 25.0, 4, (1, 3)),
            ("JC-02", "Undita pisici", "Jucarii", "Pisica", 19.0, 4, (1, 3)),
            ("JC-03", "Roata hamsteri", "Jucarii", "Rozatoare", 55.0, 1, (1, 1)),
            ("JC-04", "Oglinda papagal", "Jucarii", "Pasari", 22.0, 1, (1, 2)),
        ),
        l_nume="Nume_Magazin",
        locatii=(
            ("PetZone Bucuresti Militari", "Bucuresti", "Hipermarket", 420, 1.4),
            ("PetZone Bucuresti Floreasca", "Bucuresti", "Magazin stradal", 150, 1.0),
            ("PetZone Cluj Vivo", "Cluj-Napoca", "Mall", 180, 1.1),
            ("PetZone Timisoara Iulius", "Timisoara", "Mall", 170, 1.0),
            ("PetZone Iasi Tatarasi", "Iasi", "Magazin stradal", 130, 0.8),
            ("PetZone Brasov Coresi", "Brasov", "Mall", 160, 0.8),
            ("PetZone Ploiesti", "Ploiesti", "Magazin stradal", 110, 0.6),
            ("PetZone Arad", "Arad", "Magazin stradal", 100, 0.5),
            ("PetZone Suceava", "Suceava", "Magazin stradal", 95, 0.5),
            ("PetZone Online", "Bucuresti", "Online", None, 1.3),
        ),
        titles={"fact": ("Export vanzari · ERP PetZone", "Contine si retururi"),
                "cli": ("Clienti card PetZone", "export CRM 2025"),
                "loc": ("Retea magazine PetZone", "lista operatiuni"),
                "cal": "Calendar PetZone 2024-2025", "prod": "Catalog produse PetZone"},
    ),
    Team(
        folder="fitpro", color="#CA8A04", company="FitPro", tagline="rețea de săli de fitness",
        context=(
            "FitPro are săli de fitness în 8 orașe. Încasează abonamente, clase de grup, ședințe de personal training, spa și produse de la barul proteic.",
            "Directorul vrea un raport de 2 pagini: cât și când încasează rețeaua și ce servicii și ce membri aduc bani.",
        ),
        seed=20261104,
        fact="Incasari", t_cli="Membri", t_prod="Servicii", t_loc="Sali",
        c_id="ID_Incasare", c_data="Data", c_cli="ID_Membru", c_prod="ID_Serviciu", c_loc="ID_Sala",
        c_qty="Cantitate", c_pret="Pret_Unitar", c_canal="Canal_Plata",
        status_ok="Incasat", status_ko="Anulat",
        canale=(("Receptie", 0.45), ("Aplicatie", 0.30), ("Online", 0.15), ("Corporate", 0.10)),
        pct_ko=0.05,
        seasonality=(1.6, 1.3, 1.1, 1.0, 1.05, 0.9, 0.7, 0.7, 1.3, 1.1, 1.0, 0.8),
        cli_nume="Nume_Membru", cli_data="Data_Inscriere",
        segmente_pf=("Student", "Adult", "Senior"), segmente_pj=("Corporate",),
        firme=("SC Wellness Corp SRL", "SC Bank Fit SA", "SC Tech Active SRL", "SC Logistic Team SRL",
               "SC Health Insurance SA", "SC Media Sport SRL"),
        p_nume="Nume_Serviciu", p_cod="Cod_Serviciu", p_extra=("Durata_Min", "Număr întreg"),
        produse=(
            ("AB-01", "Abonament lunar standard", "Abonament", 60, 189.0, 12, (1, 1)),
            ("AB-02", "Abonament lunar premium", "Abonament", 90, 289.0, 7, (1, 1)),
            ("AB-03", "Abonament anual", "Abonament", 60, 1790.0, 3, (1, 1)),
            ("AB-04", "Abonament student", "Abonament", 60, 129.0, 6, (1, 1)),
            ("CG-01", "Clasa spinning", "Clase de grup", 45, 35.0, 6, (1, 4)),
            ("CG-02", "Clasa yoga", "Clase de grup", 60, 40.0, 6, (1, 4)),
            ("CG-03", "Clasa pilates", "Clase de grup", 55, 40.0, 5, (1, 4)),
            ("CG-04", "Clasa crossfit", "Clase de grup", 60, 45.0, 4, (1, 4)),
            ("PT-01", "Personal training 1 sedinta", "Personal training", 60, 120.0, 5, (1, 2)),
            ("PT-02", "Pachet PT 10 sedinte", "Personal training", 60, 990.0, 2, (1, 1)),
            ("PT-03", "Evaluare corporala", "Personal training", 30, 80.0, 3, (1, 1)),
            ("SP-01", "Sauna", "Spa", 45, 30.0, 4, (1, 3)),
            ("SP-02", "Masaj relaxare", "Spa", 50, 160.0, 3, (1, 2)),
            ("SP-03", "Masaj sportiv", "Spa", 50, 180.0, 2, (1, 2)),
            ("BR-01", "Shake proteic", "Bar", 0, 22.0, 8, (1, 3)),
            ("BR-02", "Baton proteic", "Bar", 0, 12.0, 7, (1, 5)),
            ("BR-03", "Apa minerala", "Bar", 0, 6.0, 7, (1, 4)),
            ("BR-04", "Cafea", "Bar", 0, 9.0, 5, (1, 3)),
            ("BR-05", "Creatina 300g", "Bar", 0, 119.0, 1, (1, 1)),
            ("BR-06", "Prosop FitPro", "Bar", 0, 49.0, 1, (1, 2)),
        ),
        l_nume="Nume_Sala",
        locatii=(
            ("FitPro Bucuresti Pipera", "Bucuresti", "Club mare", 2200, 1.5),
            ("FitPro Bucuresti Unirii", "Bucuresti", "Club mare", 1900, 1.3),
            ("FitPro Bucuresti Herastrau", "Bucuresti", "Boutique", 450, 0.7),
            ("FitPro Cluj Marasti", "Cluj-Napoca", "Club mare", 1700, 1.2),
            ("FitPro Timisoara", "Timisoara", "Club mic", 900, 0.8),
            ("FitPro Iasi", "Iasi", "Club mic", 850, 0.7),
            ("FitPro Brasov", "Brasov", "Club mic", 800, 0.7),
            ("FitPro Constanta", "Constanta", "Club mic", 750, 0.6),
            ("FitPro Sibiu", "Sibiu", "Boutique", 400, 0.5),
            ("FitPro Craiova", "Craiova", "Club mic", 780, 0.5),
        ),
        titles={"fact": ("Export incasari · FitPro Manager", "Nu modificati structura"),
                "cli": ("Membri FitPro", "export aplicatie membri"),
                "loc": ("Sali FitPro", "date administrative"),
                "cal": "Calendar FitPro 2024-2025", "prod": "Lista servicii FitPro"},
    ),
]


# ---------------------------------------------------------------------------
# Generarea datelor curate
# ---------------------------------------------------------------------------
def gen_calendar() -> list[dict]:
    out, d = [], CAL_START
    while d <= CAL_END:
        out.append({"Data": d, "An": d.year, "Luna": d.month, "Luna_Nume": LUNI[d.month],
                    "Trimestru": f"T{(d.month - 1) // 3 + 1}", "Zi_Saptamana": ZILE[d.weekday()],
                    "Nr_Zi": d.weekday() + 1})
        d += timedelta(days=1)
    return out


def gen_clienti(t: Team, rng: random.Random) -> list[dict]:
    orase = list(REGIUNE)
    pj_idx = list(range(3, 30, 4))[: len(t.firme)]
    pf = rng.sample(PF_NUME, 30 - len(pj_idx))
    rows = []
    for i in range(30):
        if i in pj_idx:
            nume = t.firme[pj_idx.index(i)]
            tip, seg = "PJ", rng.choice(t.segmente_pj)
        else:
            nume, tip, seg = pf.pop(), "PF", rng.choice(t.segmente_pf)
        oras = rng.choice(orase)
        d = date(2019, 1, 1) + timedelta(days=rng.randint(0, 2400))
        rows.append({t.c_cli: i + 1, t.cli_nume: nume, "Tip_Client": tip, "Oras": oras,
                     "Regiune": REGIUNE[oras], "Segment": seg, t.cli_data: d})
    return rows


def gen_produse(t: Team) -> list[dict]:
    rows = []
    for i, (cod, nume, cat, extra, pret, _pop, _q) in enumerate(t.produse, start=1):
        r = {t.c_prod: i, t.p_cod: cod, t.p_nume: nume, "Categorie": cat}
        if t.p_extra:
            r[t.p_extra[0]] = extra
        r["Pret_Lista"] = pret
        rows.append(r)
    return rows


def gen_locatii(t: Team) -> list[dict]:
    rows = []
    for i, (nume, oras, fmt, sup, _w) in enumerate(t.locatii, start=1):
        lat, lon = COORD[oras]
        rows.append({t.c_loc: i, t.l_nume: nume, "Oras": oras, "Regiune": REGIUNE[oras], "Format": fmt,
                     "Suprafata_mp": sup, "Latitudine": round(lat + (i % 3) * 0.004, 4),
                     "Longitudine": round(lon + (i % 2) * 0.005, 4)})
    return rows


def gen_fact(t: Team, rng: random.Random, n: int) -> list[dict]:
    days_by_month: dict[tuple[int, int], list[date]] = {}
    for c in gen_calendar():
        days_by_month.setdefault((c["An"], c["Luna"]), []).append(c["Data"])
    months = list(days_by_month)
    mweights = [t.seasonality[m - 1] * (1.12 if y == 2025 else 1.0) for (y, m) in months]
    pw = [p[5] for p in t.produse]
    lw = [l[4] for l in t.locatii]
    # câțiva clienți fideli; ultimii 3 nu au cumpărat nimic (DISTINCTCOUNT ≠ numărul de clienți)
    cw = [0.0 if i >= 27 else 1.0 + (3.0 if i % 7 == 0 else 0.0) for i in range(30)]
    can, canw = zip(*t.canale)
    rows = []
    for _ in range(n):
        ym = rng.choices(months, mweights)[0]
        d = rng.choice(days_by_month[ym])
        pi = rng.choices(range(len(t.produse)), pw)[0]
        prod = t.produse[pi]
        qmin, qmax = prod[6]
        q = rng.randint(qmin, qmax)
        disc = rng.choice([0, 0, 0, 0, 5, 10, 10, 15])
        val = round(q * prod[4] * (1 - disc / 100), 2)
        loc = rng.choices(range(len(t.locatii)), lw)[0] + 1
        canal = rng.choices(can, canw)[0]
        if t.locatii[loc - 1][2] == "Online" and canal == can[0]:
            canal = can[1]
        rows.append({t.c_data: d, t.c_cli: rng.choices(range(1, 31), cw)[0], t.c_prod: pi + 1,
                     t.c_loc: loc, t.c_qty: q, t.c_pret: prod[4], "Discount_Pct": disc,
                     "Valoare_Totala": val, t.c_canal: canal,
                     "Status": t.status_ko if rng.random() < t.pct_ko else t.status_ok})
    rows.sort(key=lambda r: r[t.c_data])
    out = []
    for i, r in enumerate(rows, start=1):
        out.append({t.c_id: 10000 + i, **r})
    return out


# ---------------------------------------------------------------------------
# Murdărirea · fiecare problemă are o singură rezolvare clară
# ---------------------------------------------------------------------------
def ro_date(d: date) -> str:
    return d.strftime("%d.%m.%Y")


def dirty_fact(t: Team, rows: list[dict], rng: random.Random) -> tuple[list[list], dict]:
    cols = list(rows[0])
    data, qty_txt, id_sp, st_bad, can_sp = [], 0, 0, set(), 0
    for r in rows:
        v = dict(r)
        v[t.c_data] = ro_date(r[t.c_data])
        if rng.random() < 0.12:
            v[t.c_qty] = f"{r[t.c_qty]} buc"
            qty_txt += 1
        if rng.random() < 0.06:
            v[t.c_cli] = f" {r[t.c_cli]} "
            id_sp += 1
        x = rng.random()
        if x < 0.06:
            v["Status"] = r["Status"].lower(); st_bad.add(v["Status"])
        elif x < 0.10:
            v["Status"] = r["Status"].upper(); st_bad.add(v["Status"])
        elif x < 0.13:
            v["Status"] = " " + r["Status"]; st_bad.add("«" + v["Status"] + "»")
        if rng.random() < 0.08:
            v[t.c_canal] = r[t.c_canal] + "  "
            can_sp += 1
        data.append([v[c] for c in cols])
    dup_at = sorted(rng.sample(range(10, len(data) - 10), 3), reverse=True)
    dup_ids = []
    for i in dup_at:
        data.insert(i + 1, list(data[i]))
        dup_ids.append(data[i][0])
    data.insert(len(data) // 2, [None] * len(cols))
    t1, t2 = t.titles["fact"]
    body = [[t1], [t2], cols, *data, [f"Total linii export: {len(data) - 1}"]]
    info = {"qty_txt": qty_txt, "id_sp": id_sp, "status_bad": sorted(st_bad), "can_sp": can_sp,
            "dup_ids": sorted(dup_ids), "n_dirty": len(data) - 1}
    return body, info


def dirty_clienti(t: Team, rows: list[dict], rng: random.Random) -> tuple[list[list], dict]:
    cols = list(rows[0])
    data, tip_bad, oras_sp = [], 0, 0
    for r in rows:
        v = dict(r)
        v[t.cli_data] = ro_date(r[t.cli_data])
        if rng.random() < 0.3:
            v["Tip_Client"] = "Persoana Fizica" if r["Tip_Client"] == "PF" else "Persoana Juridica"
            tip_bad += 1
        if rng.random() < 0.2:
            v["Oras"] = r["Oras"] + " "
            oras_sp += 1
        data.append([v[c] for c in cols])
    dup = rng.randrange(5, 25)
    data.append(list(data[dup]))
    t1, t2 = t.titles["cli"]
    return [[t1], [t2], cols, *data], {"tip_bad": tip_bad, "oras_sp": oras_sp, "dup_id": data[dup][0]}


def dirty_produse(t: Team, rows: list[dict], rng: random.Random) -> tuple[list[list], dict]:
    cols = list(rows[0])
    old_hdr = ["Cod", "Cod produs", "Denumire", "Cat."] + (["Detaliu"] if t.p_extra else []) + ["Pret"]
    data, prev, lei, sp = [], None, 0, 0
    for r in rows:
        v = dict(r)
        if r["Categorie"] == prev:
            v["Categorie"] = None
        prev = r["Categorie"]
        if rng.random() < 0.3:
            v["Pret_Lista"] = f"{r['Pret_Lista']:.0f} lei"
            lei += 1
        if rng.random() < 0.2:
            v[t.p_nume] = f"  {r[t.p_nume]} "
            sp += 1
        data.append([v[c] for c in cols])
    return [[t.titles["prod"]], old_hdr, cols, *data], {"lei": lei, "sp": sp}


def dirty_locatii(t: Team, rows: list[dict], rng: random.Random) -> tuple[list[list], dict]:
    cols = list(rows[0])
    data, reg_bad, mp = [], set(), 0
    for i, r in enumerate(rows):
        v = dict(r)
        if r["Regiune"] in REGIUNE_MURDAR and (i % 2 == 0):
            v["Regiune"] = REGIUNE_MURDAR[r["Regiune"]]
            reg_bad.add((v["Regiune"], r["Regiune"]))
        if r["Suprafata_mp"] is None:
            v["Suprafata_mp"] = "N/A"
        elif rng.random() < 0.4:
            v["Suprafata_mp"] = f"{r['Suprafata_mp']} mp"
            mp += 1
        v["Latitudine"] = f"{r['Latitudine']:.4f}".replace(".", ",")
        v["Longitudine"] = f"{r['Longitudine']:.4f}".replace(".", ",")
        data.append([v[c] for c in cols])
    t1, t2 = t.titles["loc"]
    has_na = any(r["Suprafata_mp"] is None for r in rows)
    return [[t1], [t2], cols, *data], {"reg_bad": sorted(reg_bad), "mp": mp, "na": has_na}


def dirty_calendar(t: Team, rows: list[dict], rng: random.Random) -> tuple[list[list], dict]:
    cols = list(rows[0])
    data, trim, low = [], 0, 0
    for r in rows:
        v = dict(r)
        v["Data"] = ro_date(r["Data"])
        if rng.random() < 0.3:
            v["Trimestru"] = "Trim " + r["Trimestru"][1]
            trim += 1
        if rng.random() < 0.1:
            v["Luna_Nume"] = r["Luna_Nume"].lower()
            low += 1
        data.append([v[c] for c in cols])
    return [[t.titles["cal"]], cols, *data], {"trim": trim, "low": low}


# ---------------------------------------------------------------------------
# Scriere Excel
# ---------------------------------------------------------------------------
def write_xlsx(path: Path, sheet: str, rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for r in rows:
        ws.append([None if c is None else c for c in r])
    for i in range(1, max(len(r) for r in rows) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def write_clean(path: Path, sheet: str, rows: list[dict]) -> None:
    cols = list(rows[0])
    write_xlsx(path, sheet, [cols, *[[r[c] for c in cols] for r in rows]])


# ---------------------------------------------------------------------------
# Construcție completă pe set de date
# ---------------------------------------------------------------------------
def build(t: Team) -> dict:
    rng = random.Random(t.seed)
    cal = gen_calendar()
    cli = gen_clienti(t, rng)
    prod = gen_produse(t)
    loc = gen_locatii(t)
    fact = gen_fact(t, rng, 236)

    d_fact, i_fact = dirty_fact(t, fact, rng)
    d_cli, i_cli = dirty_clienti(t, cli, rng)
    d_prod, i_prod = dirty_produse(t, prod, rng)
    d_loc, i_loc = dirty_locatii(t, loc, rng)
    d_cal, i_cal = dirty_calendar(t, cal, rng)

    team_dir = OUT / t.folder
    for name, rows in ((t.fact, d_fact), (t.t_cli, d_cli), (t.t_prod, d_prod), (t.t_loc, d_loc), ("Calendar", d_cal)):
        write_xlsx(team_dir / f"{name}.xlsx", name, rows)
    clean_dir = OUT_CURAT / t.folder
    for name, rows in ((t.fact, fact), (t.t_cli, cli), (t.t_prod, prod), (t.t_loc, loc), ("Calendar", cal)):
        write_clean(clean_dir / f"{name}.xlsx", name, rows)

    return {"fact": fact, "cli": cli, "prod": prod, "loc": loc, "cal": cal,
            "i_fact": i_fact, "i_cli": i_cli, "i_prod": i_prod, "i_loc": i_loc, "i_cal": i_cal}


def main() -> None:
    from proiect_html import write_team_page, write_index

    for t in TEAMS:
        data = build(t)
        write_team_page(OUT / t.folder, t, data)
        print("OK:", t.folder, len(data["fact"]), "rânduri fact")
    write_index(OUT, TEAMS)


if __name__ == "__main__":
    main()
