# -*- coding: utf-8 -*-
"""Sesiunea 11 · context în rapoarte: datele pentru widget-uri și tabelul de ținte.

Pornește de la tabelele din S10 (sesiuni/s10/date/csv/), aceleași ca în baza Azure Flyableinc.
Generează:
  sesiuni/s11/date/Tinte.xlsx, Tinte.csv   ținta de venit pe lună × companie, 2025 (= luna din 2024 × 1,12)
  js/s11-date.js                            agregate lunare pe companie și pe rută, pentru widget-urile din deck
La final afișează cifrele folosite în slide-uri și în solutii_s11.md.
Rulare: python generate_s11.py
"""
from __future__ import annotations

import csv
import json
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
COURSE = ROOT.parent.parent
SRC = COURSE / "sesiuni" / "s10" / "date" / "csv"
OUT = COURSE / "sesiuni" / "s11" / "date"
JS_OUT = COURSE / "js" / "s11-date.js"

TINTA_CRESTERE = 0.12  # ținta 2025: +12% față de aceeași lună din 2024, pentru fiecare companie
HEADER_FILL = PatternFill("solid", fgColor="0E7490")


def read(name: str) -> list[dict]:
    with (SRC / f"{name}.csv").open(encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def main() -> None:
    dates = {r["ID"]: r for r in read("Dim_Date")}
    companii = {r["ID"]: r for r in read("Dim_Companii")}
    rute = {r["ID"]: r for r in read("Dim_Rute")}
    zboruri = read("Fact_Zboruri")

    months = sorted({d["An_Luna"] for d in dates.values()})
    m_ix = {m: i for i, m in enumerate(months)}
    co_ids = sorted(companii, key=lambda i: companii[i]["Companie"])
    co_ix = {c: i for i, c in enumerate(co_ids)}
    ro_ids = sorted(rute, key=lambda i: rute[i]["Cod_Ruta"])
    ro_ix = {r: i for i, r in enumerate(ro_ids)}
    months_2425 = [m for m in months if m >= "2024-01"]
    m2_ix = {m: i for i, m in enumerate(months_2425)}

    # [venit, pasageri, locuri (zboruri neanulate), zboruri, la timp, cost, anulate]
    zero = lambda: [0.0] * 7
    by_co = [[zero() for _ in months] for _ in co_ids]
    by_ro = [[zero() for _ in months_2425] for _ in ro_ids]
    for z in zboruri:
        d = dates[z["ID_Date"]]
        venit = float(z["Venit_Bilete"]) + float(z["Venit_Auxiliar"])
        anulat = z["Status_Zbor"] == "Anulat"
        row = [venit, int(z["Pasageri_Imbarcati"]), 0 if anulat else int(z["Locuri_Disponibile"]), 1,
               1 if z["Status_Zbor"] == "La timp" else 0, float(z["Cost_Combustibil"]) + float(z["Cost_Operare"]), 1 if anulat else 0]
        cells = [by_co[co_ix[z["ID_Companie"]]][m_ix[d["An_Luna"]]]]
        if d["An_Luna"] in m2_ix:
            cells.append(by_ro[ro_ix[z["ID_Ruta"]]][m2_ix[d["An_Luna"]]])
        for cell in cells:
            for k, v in enumerate(row):
                cell[k] += v

    def pack(rows):
        return [[[round(v) for v in cell] for cell in series] for series in rows]

    # Ținte: prima zi a lunii (ID_Date din Dim_Date) × companie
    first_day = {d["Data"][:7]: d for d in dates.values() if d["Data"].endswith("-01")}
    tinte = []
    for cid in co_ids:
        for m in [x for x in months if x.startswith("2025")]:
            py = by_co[co_ix[cid]][m_ix["2024" + m[4:]]][0]
            fd = first_day[m]
            tinte.append({"ID_Date": int(fd["ID"]), "Data": date.fromisoformat(fd["Data"]), "ID_Companie": int(cid),
                          "Tinta_Venit": int(round(py * (1 + TINTA_CRESTERE), -2))})
    tinte.sort(key=lambda r: (r["ID_Date"], r["ID_Companie"]))
    write_tinte(tinte)

    payload = {
        "months": months,
        "companii": [companii[c]["Companie"] for c in co_ids],
        "tipCompanie": [companii[c]["Tip_Companie"] for c in co_ids],
        "co": pack(by_co),
        "monthsRute": months_2425,
        "rute": [rute[r]["Cod_Ruta"] for r in ro_ids],
        "tipRuta": [rute[r]["Tip_Ruta"] for r in ro_ids],
        "ro": pack(by_ro),
        "tinte": [[t["Tinta_Venit"] for t in tinte if t["ID_Companie"] == int(c)] for c in co_ids],
        "tintaCrestere": TINTA_CRESTERE,
    }
    JS_OUT.write_text("/* Generat de assets/s11/generate_s11.py — nu editați de mână.\n"
                      "   co[companie][lună] și ro[rută][lună 2024–2025] = [venit, pasageri, locuri, zboruri, la timp, cost, anulate]. */\n"
                      "window.S11_DATE = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n", encoding="utf-8")
    report(payload, months)


def write_tinte(rows: list[dict]) -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    cols = list(rows[0])
    with (OUT / "Tinte.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(cols)
        for r in rows:
            w.writerow([r[c].isoformat() if isinstance(r[c], date) else r[c] for c in cols])
    wb = Workbook()
    ws = wb.active
    ws.title = "Tinte"
    ws.append(cols)
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(c) + 4)
    for rr in range(2, len(rows) + 2):
        ws.cell(row=rr, column=2).number_format = "DD.MM.YYYY"
        ws.cell(row=rr, column=4).number_format = "#,##0"
    ws.freeze_panes = "A2"
    wb.properties.created = wb.properties.modified = __import__("datetime").datetime(2026, 1, 1)
    wb.save(OUT / "Tinte.xlsx")


def report(p: dict, months: list[str]) -> None:
    """Cifrele din slide-uri: totaluri pe an, luni, trimestre, companii."""
    def tot(pred, k=0, src="co"):
        ms = p["months"] if src == "co" else p["monthsRute"]
        return sum(cell[k] for series in p[src] for m, cell in zip(ms, series) if pred(m))

    print("Tinte: %d rânduri" % sum(len(t) for t in p["tinte"]))
    for an in ("2023", "2024", "2025"):
        v = tot(lambda m: m.startswith(an))
        lf = tot(lambda m: m.startswith(an), 1) / tot(lambda m: m.startswith(an), 2)
        ot = tot(lambda m: m.startswith(an), 4) / tot(lambda m: m.startswith(an), 3)
        print(f"{an}: venit {v:,.0f}  pasageri {tot(lambda m: m.startswith(an), 1):,.0f}  LF {lf:.2%}  la timp {ot:.2%}")
    print("Ținta 2025:", f"{sum(map(sum, p['tinte'])):,.0f}")
    for mm in range(1, 13):
        a = tot(lambda m: m == f"2025-{mm:02d}")
        b = tot(lambda m: m == f"2024-{mm:02d}")
        print(f"luna {mm:02d}: 2025 {a:,.0f}  2024 {b:,.0f}  YoY {a / b - 1:+.1%}")
    q = lambda an, t: tot(lambda m: m.startswith(an) and (int(m[5:]) - 1) // 3 + 1 == t)
    for an in ("2024", "2025"):
        print(an, "trimestre:", [f"{q(an, t):,.0f}" for t in range(1, 5)])
    for i, c in enumerate(p["companii"]):
        a = sum(cell[0] for m, cell in zip(months, p["co"][i]) if m.startswith("2025"))
        b = sum(cell[0] for m, cell in zip(months, p["co"][i]) if m.startswith("2024"))
        print(f"{c:18s} 2025 {a:12,.0f}  YoY {a / b - 1:+.1%}  vs țintă {a / sum(p['tinte'][i]) - 1:+.1%}")


if __name__ == "__main__":
    main()
