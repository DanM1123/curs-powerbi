# -*- coding: utf-8 -*-
"""Generează seturi murdare + brief proiect pentru echipele S5 (rosu, albastru, verde, galben)."""
from __future__ import annotations

import random
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path
from textwrap import dedent

from openpyxl import Workbook

import generate_s05_datasets as g5

ROOT = Path(__file__).resolve().parent
SESSION_S05 = ROOT.parent.parent / "sesiuni" / "s05"
STATUS = ["Platit", "Partial", "Neplatit", "Anulat"]
STATUS_MURDAR = ["Platit", "Platit ", "PLATIT", "Partial", "partial", "Neplatit", "Neplatit ", "Anulat", "anulat"]


@dataclass(frozen=True)
class TeamSpec:
    folder: str
    label: str
    members: tuple[str, ...]
    color: str
    seed: int
    company: str
    tagline: str
    industry: str
    context: str
    context_detail: str
    fact_title: str
    fact_sub: str
    mag_sheet_title: str
    client_j1: str
    client_j2: str
    dup_client_name: str
    magazine: tuple
    produse: tuple
    clienti: tuple
    canal: tuple
    online_mag_id: int
    pages: tuple  # (title, focus, questions, (rec1, rec2))


def _mag(*rows):
    return rows


def _prod(*rows):
    return rows


def _cli(*rows):
    return rows


TEAMS: list[TeamSpec] = [
    TeamSpec(
        folder="echipa_rosu",
        label="Echipa Roșu",
        members=(
            "Alina Matasaru",
            "Andreea Filip",
            "Karola Mocan",
            "Alexandru-Paul Dima",
            "Edgar Denes",
        ),
        color="#DC2626",
        seed=20260401,
        company="GustRo",
        tagline="lanț de restaurante & catering",
        industry="HoReCa",
        context=(
            "GustRo are restaurante în mai multe orașe. Vinde mâncare și băuturi la locație, la livrare sau catering."
        ),
        context_detail=(
            "În Excel: linii de vânzare (produs, cantitate, valoare, restaurant, client, dată, canal) "
            "și tabele Dim pentru clienți, meniu, locații și calendar. "
            "Raportul e pentru managementul GustRo: înțelege performanța rețelei, nu doar un total."
        ),
        fact_title="Export comenzi · POS GustRo v2",
        fact_sub="Nu modifica — sincronizare noaptea",
        mag_sheet_title="Locații restaurant",
        client_j1="CRM clienți · GustRo",
        client_j2="export Q1-2025",
        dup_client_name="Maria Ionescu",
        magazine=_mag(
            ("GR001", "GustRo Cluj Old Town", "Cluj-Napoca", 46.769, 23.589, "Restaurant", 180),
            ("GR002", "GustRo Unirii București", "Bucuresti", 44.426, 26.102, "Restaurant", 220),
            ("GR003", "GustRo Timișoara", "Timisoara", 45.749, 21.209, "Restaurant", 160),
            ("GR004", "GustRo Mamaia", "Constanta", 44.250, 28.632, "Sezonier", 140),
            ("GR005", "GustRo Iasi Palas", "Iasi", 47.158, 27.601, "Restaurant", 150),
            ("GR006", "GustRo Brasov Centru", "Brasov", 45.658, 25.601, "Restaurant", 170),
            ("GR007", "GustRo Oradea", "Oradea", 47.046, 21.919, "Franciza", 120),
            ("GR008", "GustRo Craiova", "Craiova", 44.330, 23.795, "Restaurant", 130),
            ("GR009", "Kitchen Delivery Hub", "Bucuresti", 44.438, 26.098, "Delivery hub", 90),
            ("GR010", "Food truck Galati", "Galati", 45.435, 28.008, "Food truck", 45),
        ),
        produse=_prod(
            ("MN-E01", "Burger GustRo", "Mancare", "Fel principal", "Bucatarie", 42.0),
            ("MN-E02", "Paste carbonara", "Mancare", "Fel principal", "Bucatarie", 38.0),
            ("MN-E03", "Salata Caesar", "Mancare", "Aperitiv", "Bucatarie", 28.0),
            ("MN-E04", "Ciorba radauteana", "Mancare", "Supe", "Bucatarie", 22.0),
            ("MN-E05", "Pizza Margherita", "Mancare", "Pizza", "Bucatarie", 35.0),
            ("MN-B01", "Limonada casa", "Bauturi", "Non-alcool", "Bar", 18.0),
            ("MN-B02", "Espresso", "Bauturi", "Cafea", "Bar", 9.0),
            ("MN-B03", "Vin rosu pahar", "Bauturi", "Alcool", "Bar", 24.0),
            ("MN-B04", "Apa minerala 0.5L", "Bauturi", "Non-alcool", "Bar", 8.0),
            ("MN-D01", "Cheesecake", "Desert", "Torturi", "Patiserie", 26.0),
            ("MN-D02", "Papanasi", "Desert", "Traditional", "Patiserie", 29.0),
            ("MN-C01", "Meniu copii", "Mancare", "Copii", "Bucatarie", 32.0),
            ("MN-C02", "Platou catering 10 pers", "Catering", "Evenimente", "Catering", 320.0),
            ("MN-C03", "Coffee break office", "Catering", "Corporate", "Catering", 450.0),
            ("MN-E06", "Tacos pui", "Mancare", "Street food", "Bucatarie", 36.0),
            ("MN-B05", "Bere draft 0.5L", "Bauturi", "Alcool", "Bar", 16.0),
            ("MN-D03", "Inghetata artizanala", "Desert", "Inghetata", "Patiserie", 15.0),
            ("MN-E07", "Wrap vegan", "Mancare", "Fel principal", "Bucatarie", 34.0),
        ),
        clienti=_cli(
            ("CR001", "SC Event Masters SRL", "PJ", "Cluj-Napoca", "Corporate", 8.9),
            ("CR002", "SC Office Lunch SA", "PJ", "Bucuresti", "Corporate", 8.2),
            ("CR003", "Ion Dumitrescu", "PF", "Timisoara", "Familie", 7.5),
            ("CR004", "SC Hotel Delta SRL", "PJ", "Constanta", "HoReCa partener", 8.0),
            ("CR005", "Maria Ionescu", "PF", "Brasov", "Familie", 9.1),
            ("CR006", "SC Uni Campus SRL", "PJ", "Iasi", "Corporate", 7.8),
            ("CR007", "SC Tech Park Food SRL", "PJ", "Oradea", "Corporate", 8.4),
            ("CR008", "Andrei Pop", "PF", "Sibiu", "Familie", 6.9),
            ("CR009", "SC Wedding Plan SRL", "PJ", "Craiova", "Evenimente", 8.6),
            ("CR010", "Elena Radu", "PF", "Galati", "Familie", 8.3),
            ("CR011", "SC Airline Catering SRL", "PJ", "Bucuresti", "Corporate", 7.7),
            ("CR012", "SC Mall Food Court SRL", "PJ", "Ploiesti", "Retail partener", 7.4),
            ("CR013", "Catalin Moldovan", "PF", "Braila", "Familie", 7.1),
            ("CR014", "SC Pharma Events SRL", "PJ", "Pitesti", "Corporate", 8.1),
            ("CR015", "Diana Luca", "PF", "Voluntari", "Familie", 8.7),
            ("CR016", "SC Sport Club SRL", "PJ", "Targu Mures", "Corporate", 7.9),
            ("CR017", "Radu Enache", "PF", "Alba Iulia", "Familie", 7.6),
            ("CR018", "SC Festival Stage SRL", "PJ", "Baia Mare", "Evenimente", 8.0),
            ("CR019", "Simona Nita", "PF", "Pantelimon", "Familie", 9.0),
            ("CR020", "SC Startup Hub SRL", "PJ", "Cluj-Napoca", "Corporate", 8.5),
            ("CR021", "Victor Marinescu", "PF", "Oradea", "Familie", 6.8),
            ("CR022", "SC Beauty Expo SRL", "PJ", "Bucuresti", "Evenimente", 8.3),
            ("CR023", "Ana Cojocaru", "PF", "Iasi", "Familie", 7.4),
            ("CR024", "SC Port Services SRL", "PJ", "Constanta", "Corporate", 7.6),
            ("CR025", "George Ilie", "PF", "Craiova", "Familie", 8.8),
            ("CR026", "SC Mountain Resort SRL", "PJ", "Brasov", "HoReCa partener", 8.2),
            ("CR027", "Laura Sandu", "PF", "Timisoara", "Familie", 8.0),
            ("CR028", "SC Logistics Lunch SRL", "PJ", "Sibiu", "Corporate", 7.3),
        ),
        canal=("La masa", "Livrare", "Glovo", "Telefon comanda"),
        online_mag_id=9,
        pages=(
            (
                "Performanță vânzări",
                "Prima pagină = imagine de ansamblu pentru directorul de operațiuni: cifre mari + trend + locații.",
                (
                    "Care este valoarea totală a vânzărilor (ex. sumă Valoare_Totala)?",
                    "Cum se schimbă vânzările de la o lună la alta sau pe ani?",
                    "Există diferențe clare între regiuni sau orașe?",
                    "Care 5 restaurante sau locații ies în față?",
                ),
                (
                    "Recomandare: 2–3 carduri (total, număr linii, medie) — tu alegi ce are sens.",
                    "Recomandare: linie pe calendar (lună/an) și coloane pe locație sau regiune.",
                ),
            ),
            (
                "Mix meniu",
                "Pagina pentru bucătar-șef și marketing: ce categorii și produse trag vânzarea.",
                (
                    "Cum se împarte valoarea între Mancare, Bauturi, Desert și Catering?",
                    "Ce subcategorii (ex. Fel principal, Bar) sunt puternice?",
                    "Care produse din meniu sunt cele mai vândute ca valoare?",
                    "Catering-ul (evenimente, corporate) contează mult față de restaurant?",
                ),
                (
                    "Recomandare: diagramă stacked sau donut pe categorie.",
                    "Recomandare: bar chart sortat pe produse sau subcategorii.",
                ),
            ),
            (
                "Clienți",
                "Echipa comercială vrea să vadă cine aduce valoare: familii, corporate, evenimente.",
                (
                    "Ce pondere au clienții PF față de PJ?",
                    "Ce segmente (Corporate, Familie, Evenimente…) domină?",
                    "Cine sunt clienții cu cele mai mari comenzi?",
                    "Există orașe unde un tip de client e mai prezent?",
                ),
                (
                    "Recomandare: donut PF/PJ sau pe Segment din Dim_Clienti.",
                    "Recomandare: tabel sau bar chart cu top clienți + coloană valoare.",
                ),
            ),
            (
                "Locații și canale",
                "Unde investim (locație, delivery) și cum comandă oamenii.",
                (
                    "Ce formate (restaurant, delivery hub, food truck) se comportă diferit?",
                    "Compară vânzările pe canal: La masa, Livrare, Glovo, Telefon comanda.",
                    "Există locații unde delivery-ul e mult mai mare decât la masă?",
                ),
                (
                    "Recomandare: bar chart pe Nume_Magazin sau Format_Magazin.",
                    "Recomandare: coloane pe Canal_Vanzare; opțional matrice locație × canal.",
                ),
            ),
        ),
    ),
    TeamSpec(
        folder="echipa_albastru",
        label="Echipa Albastru",
        members=(
            "Violeta Gutu",
            "Patricia Cadar",
            "Costin Văetiși",
            "Ovidiu Borlean",
        ),
        color="#2563EB",
        seed=20260402,
        company="FixAuto",
        tagline="rețea service auto & piese",
        industry="Automotive aftermarket",
        context=(
            "FixAuto are service-uri auto în mai multe orașe: revizii, ITP, piese și manoperă."
        ),
        context_detail=(
            "Fiecare linie = o poziție pe deviz (manoperă sau piesă), cu centru service, client și dată. "
            "Dimensiunile descriu clienți (inclusiv flote), catalog servicii/piese, centre și calendar. "
            "Raportul sprijină decizii despre încărcarea centrelor și mix servicii vs piese."
        ),
        fact_title="Export linii service · DMS FixAuto",
        fact_sub="Confidential — uz intern",
        mag_sheet_title="Centre service",
        client_j1="Baza clienti DMS",
        client_j2="backup martie 2025",
        dup_client_name="Ion Dumitrescu",
        magazine=_mag(
            ("FA001", "FixAuto Cluj Floresti", "Cluj-Napoca", 46.771, 23.623, "Service complet", 850),
            ("FA002", "FixAuto Bucuresti Baneasa", "Bucuresti", 44.503, 26.078, "Service complet", 1200),
            ("FA003", "FixAuto Timisoara", "Timisoara", 45.749, 21.209, "Service + ITP", 720),
            ("FA004", "FixAuto Constanta", "Constanta", 44.160, 28.635, "Service complet", 680),
            ("FA005", "FixAuto Iasi", "Iasi", 47.159, 27.602, "Service + anvelope", 640),
            ("FA006", "FixAuto Brasov", "Brasov", 45.658, 25.601, "Service complet", 700),
            ("FA007", "FixAuto Oradea", "Oradea", 47.047, 21.919, "Franciza", 520),
            ("FA008", "FixAuto Craiova", "Craiova", 44.330, 23.795, "Service + ITP", 580),
            ("FA009", "FixAuto Express Online", "Bucuresti", 44.437, 26.097, "Programari online", 0),
            ("FA010", "FixAuto Pitesti", "Pitesti", 44.856, 24.869, "Service mic", 410),
        ),
        produse=_prod(
            ("SV-01", "Revizie periodica", "Servicii", "Mentenanta", "FixAuto", 349.0),
            ("SV-02", "ITP autoturism", "Servicii", "ITP", "FixAuto", 145.0),
            ("SV-03", "Schimb ulei + filtre", "Servicii", "Mentenanta", "FixAuto", 279.0),
            ("SV-04", "Diagnoza computerizata", "Servicii", "Diagnoza", "FixAuto", 189.0),
            ("SV-05", "Geometrie roti", "Servicii", "Anvelope", "FixAuto", 220.0),
            ("PZ-01", "Anvelopa vara 205/55", "Piese", "Anvelope", "Michelin", 420.0),
            ("PZ-02", "Set placute frana", "Piese", "Frane", "Bosch", 310.0),
            ("PZ-03", "Baterie auto 70Ah", "Piese", "Electrice", "Varta", 580.0),
            ("PZ-04", "Ulei motor 5W30 5L", "Piese", "Lubrifianti", "Castrol", 199.0),
            ("PZ-05", "Filtru aer", "Piese", "Filtre", "Mann", 89.0),
            ("SV-06", "Inlocuire ambreiaj", "Servicii", "Transmisie", "FixAuto", 1890.0),
            ("PZ-06", "Discuri frana fata", "Piese", "Frane", "TRW", 450.0),
            ("SV-07", "Contract flota revizii", "Servicii", "Flota", "FixAuto", 990.0),
            ("PZ-07", "Lampa LED H7", "Piese", "Iluminat", "Philips", 65.0),
            ("SV-08", "Montaj anvelope", "Servicii", "Anvelope", "FixAuto", 120.0),
            ("PZ-08", "Antigel 5L", "Piese", "Lubrifianti", "Total", 95.0),
            ("SV-09", "Inspectie pre-ITP", "Servicii", "ITP", "FixAuto", 99.0),
            ("PZ-09", "Stergatoare parbriz set", "Piese", "Accesorii", "Bosch", 78.0),
        ),
        clienti=_cli(
            ("CA001", "SC Flota Rapid SRL", "PJ", "Bucuresti", "Flota", 8.7),
            ("CA002", "SC Taxi Union SA", "PJ", "Cluj-Napoca", "Flota", 7.9),
            ("CA003", "Ion Dumitrescu", "PF", "Timisoara", "Retail", 8.1),
            ("CA004", "SC Rent A Car Delta SRL", "PJ", "Constanta", "Flota", 8.4),
            ("CA005", "Maria Ionescu", "PF", "Brasov", "Retail", 9.0),
            ("CA006", "SC Courier Express SRL", "PJ", "Iasi", "Flota", 7.6),
            ("CA007", "SC Construct Utilaje SRL", "PJ", "Oradea", "Flota", 8.2),
            ("CA008", "Andrei Pop", "PF", "Sibiu", "Retail", 7.0),
            ("CA009", "SC Ambulance Service SRL", "PJ", "Craiova", "Flota", 8.8),
            ("CA010", "Elena Radu", "PF", "Galati", "Retail", 8.3),
            ("CA011", "SC City Bus SRL", "PJ", "Ploiesti", "Flota", 7.5),
            ("CA012", "SC Premium Leasing SA", "PJ", "Bucuresti", "Flota", 8.9),
            ("CA013", "Catalin Moldovan", "PF", "Braila", "Retail", 6.9),
            ("CA014", "SC Agro Machines SRL", "PJ", "Pitesti", "Flota", 7.8),
            ("CA015", "Diana Luca", "PF", "Voluntari", "Retail", 8.6),
            ("CA016", "SC Police Fleet SRL", "PJ", "Targu Mures", "Institutional", 8.0),
            ("CA017", "Radu Enache", "PF", "Alba Iulia", "Retail", 7.4),
            ("CA018", "SC Mining Trucks SRL", "PJ", "Baia Mare", "Flota", 7.2),
            ("CA019", "Simona Nita", "PF", "Pantelimon", "Retail", 9.1),
            ("CA020", "SC IT Company Cars SRL", "PJ", "Cluj-Napoca", "Flota", 8.5),
            ("CA021", "Victor Marinescu", "PF", "Oradea", "Retail", 6.7),
            ("CA022", "SC Hotel Shuttle SRL", "PJ", "Bucuresti", "Flota", 8.1),
            ("CA023", "Ana Cojocaru", "PF", "Iasi", "Retail", 7.5),
            ("CA024", "SC Port Logistics SRL", "PJ", "Constanta", "Flota", 8.0),
            ("CA025", "George Ilie", "PF", "Craiova", "Retail", 8.7),
            ("CA026", "SC Ski Resort Transport SRL", "PJ", "Brasov", "Flota", 7.9),
            ("CA027", "Laura Sandu", "PF", "Timisoara", "Retail", 8.2),
            ("CA028", "SC School Bus SRL", "PJ", "Sibiu", "Institutional", 8.3),
        ),
        canal=("Programare online", "Walk-in", "Flota B2B", "Telefon"),
        online_mag_id=9,
        pages=(
            (
                "Performanță rețea",
                "Managementul FixAuto vrea total încasări, evoluție și centre care trac.",
                (
                    "Care este valoarea totală a lucrărilor facturate?",
                    "Cum arată trendul veniturilor pe luni sau trimestre?",
                    "Ce regiuni sau orașe generează cele mai multe încasări?",
                    "Top centre service după valoare?",
                ),
                (
                    "Recomandare: carduri KPI + diagramă linie pe Dim_Date.",
                    "Recomandare: coloane pe Regiune sau Nume_Magazin (centru).",
                ),
            ),
            (
                "Servicii vs piese",
                "Atelier și stoc: manoperă sau vânzare de piese domină?",
                (
                    "Ce procent din valoare vine din categoria Servicii față de Piese?",
                    "Ce tipuri de lucrări (Mentenanta, ITP, Anvelope…) sunt cerute cel mai des?",
                    "Care produse/servicii (SKU) aduc cele mai multe încasări?",
                ),
                (
                    "Recomandare: donut sau coloane pe Categorie din Dim_Produse.",
                    "Recomandare: bar chart pe Subcategorie sau Nume_Produs.",
                ),
            ),
            (
                "Clienți și flote",
                "Retail (șofer individual) vs contracte B2B și flote.",
                (
                    "Cum se compară PF cu PJ ca valoare totală?",
                    "Segmentele Flota, Retail, Institutional — care contează mai mult?",
                    "Care clienți PJ (flote) sunt în top?",
                ),
                (
                    "Recomandare: diagramă Tip_Client sau Segment.",
                    "Recomandare: tabel/top N clienți cu Denumire și valoare.",
                ),
            ),
            (
                "Centre și acces",
                "Capacitate și marketing: cum ajung mașinile în service.",
                (
                    "Centrele au același ritm de încasări sau există lideri?",
                    "Programare online vs walk-in vs Flota B2B — cum se împarte valoarea?",
                    "Unde ar merita extindere (centru puternic vs slab)?",
                ),
                (
                    "Recomandare: bar chart pe centru; opțional pe Format_Magazin.",
                    "Recomandare: coloane sau donut pe Canal_Vanzare.",
                ),
            ),
        ),
    ),
    TeamSpec(
        folder="echipa_verde",
        label="Echipa Verde",
        members=(
            "Silvana Boboescu",
            "Larisa Pana",
            "Adrian Marginea",
            "Bogdan Bodescu",
        ),
        color="#059669",
        seed=20260403,
        company="FarmPlus",
        tagline="rețea farmacii",
        industry="Retail farmaceutic",
        context=(
            "FarmPlus este o rețea de farmacii. Vinde medicamente OTC, cosmetice, suplimente și alte produse."
        ),
        context_detail=(
            "Linii de vânzare din casă și online: produs, cantitate, valoare, farmacie, client, canal. "
            "Dimensiuni pentru clienți (fidelity, walk-in, contracte), produse (OTC, RX, cosmetice…), farmacii, calendar. "
            "Regional managerii compară farmaciile și categoriile de asortiment."
        ),
        fact_title="Export vanzari · ERP FarmPlus",
        fact_sub="Include retururi marcate Anulat",
        mag_sheet_title="Farmacii",
        client_j1="Clienti fidelity + PJ",
        client_j2="GDPR anonim partial",
        dup_client_name="Elena Radu",
        magazine=_mag(
            ("FP001", "FarmPlus Cluj Marasti", "Cluj-Napoca", 46.771, 23.623, "Farmacie strada", 95),
            ("FP002", "FarmPlus Unirii", "Bucuresti", 44.426, 26.102, "Farmacie flagship", 140),
            ("FP003", "FarmPlus Timisoara", "Timisoara", 45.749, 21.209, "Farmacie mall", 110),
            ("FP004", "FarmPlus Constanta", "Constanta", 44.160, 28.635, "Farmacie strada", 88),
            ("FP005", "FarmPlus Iasi", "Iasi", 47.159, 27.602, "Farmacie spital", 102),
            ("FP006", "FarmPlus Brasov", "Brasov", 45.658, 25.601, "Farmacie strada", 92),
            ("FP007", "FarmPlus Oradea", "Oradea", 47.047, 21.919, "Farmacie franciza", 78),
            ("FP008", "FarmPlus Craiova", "Craiova", 44.330, 23.795, "Farmacie strada", 85),
            ("FP009", "FarmPlus E-com", "Bucuresti", 44.437, 26.097, "Online", 0),
            ("FP010", "FarmPlus Ploiesti", "Ploiesti", 44.940, 26.022, "Farmacie strada", 70),
        ),
        produse=_prod(
            ("OTC-01", "Paracetamol 500mg", "OTC", "Analgezice", "PharmaRo", 12.0),
            ("OTC-02", "Vitamina C eff", "OTC", "Vitamine", "Health+", 28.0),
            ("OTC-03", "Sirop tuse copii", "OTC", "Respirator", "KidCare", 32.0),
            ("RX-01", "Antibiotic rx pack", "RX", "Antibiotice", "PharmaRo", 89.0),
            ("RX-02", "Antihipertensiv rx", "RX", "Cardio", "CardioMed", 65.0),
            ("COS-01", "Crema fata SPF50", "Cosmetice", "Dermato", "DermaLux", 79.0),
            ("COS-02", "Sampon dermatologic", "Cosmetice", "Par", "DermaLux", 54.0),
            ("SUP-01", "Omega 3 capsule", "Suplimente", "Nutritie", "NutriMax", 95.0),
            ("SUP-02", "Probiotice 30cps", "Suplimente", "Digestie", "NutriMax", 110.0),
            ("OTC-04", "Plasturi assort", "OTC", "Prim ajutor", "MedKit", 18.0),
            ("OTC-05", "Test antigen", "OTC", "Diagnostic", "QuickTest", 25.0),
            ("COS-03", "Apa micelara 400ml", "Cosmetice", "Demachiere", "PureSkin", 42.0),
            ("RX-03", "Insulina rx", "RX", "Diabet", "DiabetCare", 220.0),
            ("OTC-06", "Gel maini antibacterian", "OTC", "Igiena", "CleanHands", 15.0),
            ("SUP-03", "Magnesium B6", "Suplimente", "Nutritie", "NutriMax", 48.0),
            ("COS-04", "Ruj mat nude", "Cosmetice", "Make-up", "GlamUP", 59.0),
            ("OTC-07", "Bandaj elastic", "OTC", "Prim ajutor", "MedKit", 22.0),
            ("RX-04", "Stabilizator glicemie rx", "RX", "Diabet", "DiabetCare", 145.0),
        ),
        clienti=_cli(
            ("FV001", "SC Clinica Regina SRL", "PJ", "Bucuresti", "Institutional", 8.5),
            ("FV002", "SC Spital Judetean SA", "PJ", "Cluj-Napoca", "Institutional", 8.8),
            ("FV003", "Ion Popescu", "PF", "Timisoara", "Fidelity", 7.8),
            ("FV004", "SC Casa de nursing SRL", "PJ", "Brasov", "Institutional", 8.0),
            ("FV005", "Elena Radu", "PF", "Constanta", "Fidelity", 9.0),
            ("FV006", "SC Cabinet medicina muncii SRL", "PJ", "Iasi", "Corporate", 7.7),
            ("FV007", "SC Gym Chain SRL", "PJ", "Oradea", "Corporate", 8.1),
            ("FV008", "Andrei Stan", "PF", "Sibiu", "Walk-in", 6.8),
            ("FV009", "SC Dental Clinic SRL", "PJ", "Craiova", "Corporate", 8.3),
            ("FV010", "Maria Ionescu", "PF", "Galati", "Fidelity", 8.4),
            ("FV011", "SC School Health SRL", "PJ", "Ploiesti", "Institutional", 7.9),
            ("FV012", "SC Beauty Salon Chain SRL", "PJ", "Bucuresti", "Corporate", 8.2),
            ("FV013", "Catalin Moldovan", "PF", "Braila", "Walk-in", 7.2),
            ("FV014", "SC Pharma Distributor SRL", "PJ", "Pitesti", "B2B", 8.6),
            ("FV015", "Diana Luca", "PF", "Voluntari", "Fidelity", 8.9),
            ("FV016", "SC Vet Clinic SRL", "PJ", "Targu Mures", "Corporate", 7.5),
            ("FV017", "Radu Enache", "PF", "Alba Iulia", "Walk-in", 7.6),
            ("FV018", "SC Factory Medical SRL", "PJ", "Baia Mare", "Corporate", 8.0),
            ("FV019", "Simona Nita", "PF", "Pantelimon", "Fidelity", 9.2),
            ("FV020", "SC University Campus SRL", "PJ", "Cluj-Napoca", "Institutional", 8.1),
            ("FV021", "Victor Marinescu", "PF", "Oradea", "Walk-in", 6.9),
            ("FV022", "SC Elder Care Home SRL", "PJ", "Bucuresti", "Institutional", 8.7),
            ("FV023", "Ana Cojocaru", "PF", "Iasi", "Fidelity", 7.4),
            ("FV024", "SC Hotel Medical SRL", "PJ", "Constanta", "Corporate", 7.8),
            ("FV025", "George Ilie", "PF", "Craiova", "Walk-in", 8.6),
            ("FV026", "SC Ski Clinic SRL", "PJ", "Brasov", "Corporate", 8.0),
            ("FV027", "Laura Sandu", "PF", "Timisoara", "Fidelity", 8.1),
            ("FV028", "SC Insurance Check SRL", "PJ", "Sibiu", "Corporate", 7.6),
        ),
        canal=("Farmacie", "Comanda online", "Contract PJ", "Telefon"),
        online_mag_id=9,
        pages=(
            (
                "Performanță rețea",
                "Prima pagină: cifra de afaceri totală, evoluție, farmacii de top.",
                (
                    "Care este valoarea totală a vânzărilor?",
                    "Există creștere sau scăderi pe parcursul lunilor din date?",
                    "Ce regiuni sau orașe au farmacii mai performante?",
                    "Care farmacii (inclusiv flagship vs mall) ies în top?",
                ),
                (
                    "Recomandare: carduri + linie temporală din Dim_Date.",
                    "Recomandare: bar chart pe Nume_Magazin sau Regiune.",
                ),
            ),
            (
                "Asortiment",
                "Merchandising: OTC, RX, cosmetice, suplimente.",
                (
                    "Cum se împarte valoarea pe categorii principale?",
                    "RX vs restul — ce pondere are (doar ca valoare, fără detaliu medical)?",
                    "Care produse individuale sunt cele mai vândute?",
                    "Cosmeticele și suplimentele cresc față de OTC?",
                ),
                (
                    "Recomandare: donut/coloane pe Categorie.",
                    "Recomandare: top produse — bar chart pe Nume_Produs.",
                ),
            ),
            (
                "Clienți",
                "Card fidelity, walk-in, contracte cu firme sau instituții.",
                (
                    "PF vs PJ — cine aduce mai multă valoare?",
                    "Segmente (Fidelity, Walk-in, Institutional, Corporate): care domină?",
                    "Clienții PJ (clinici, firme) — cine e în top?",
                ),
                (
                    "Recomandare: diagramă Tip_Client + eventual Segment.",
                    "Recomandare: tabel clienți cu oraș și valoare agregată.",
                ),
            ),
            (
                "Farmacii și canale",
                "Rețea fizică vs comenzi online.",
                (
                    "Farmaciile au profiluri diferite de vânzare?",
                    "Canal Farmacie vs Comanda online — comparație valoare?",
                    "Unde ar prioritiza managerul investiția (locație sau digital)?",
                ),
                (
                    "Recomandare: bar chart pe farmacie; filtru opțional pe Format.",
                    "Recomandare: coloane pe Canal_Vanzare din fact.",
                ),
            ),
        ),
    ),
    TeamSpec(
        folder="echipa_galben",
        label="Echipa Galben",
        members=(
            "Mariana Vasile",
            "Beatrice Antonievici",
            "Ruxandra Gheorghita",
            "Horea Stanescu",
        ),
        color="#CA8A04",
        seed=20260404,
        company="BuildMat",
        tagline="depozite materiale construcții",
        industry="Construcții & bricolaj",
        context=(
            "BuildMat vinde materiale de construcții și unelte din depozite din țară."
        ),
        context_detail=(
            "Linii de vânzare din depozite: material sau unealtă, cantitate, valoare, client, canal (depozit, livrare șantier, B2B). "
            "Dimensiuni pentru constructori vs bricolaj, categorii produse, depozite și calendar. "
            "Conducerea urmărește volumul pe regiuni și mix materiale vs unelte."
        ),
        fact_title="Export linii vanzari · WMS BuildMat",
        fact_sub="Include comenzi partial livrate",
        mag_sheet_title="Depozite materiale",
        client_j1="Clienti CRM BuildMat",
        client_j2="export februarie 2025",
        dup_client_name="George Ilie",
        magazine=_mag(
            ("BM001", "BuildMat Cluj Apahida", "Cluj-Napoca", 46.771, 23.623, "Depozit mare", 4200),
            ("BM002", "BuildMat Bucuresti Sud", "Bucuresti", 44.380, 26.120, "Depozit mare", 5800),
            ("BM003", "BuildMat Timisoara", "Timisoara", 45.749, 21.209, "Depozit", 3100),
            ("BM004", "BuildMat Constanta", "Constanta", 44.160, 28.635, "Depozit", 2900),
            ("BM005", "BuildMat Iasi", "Iasi", 47.159, 27.602, "Depozit", 2700),
            ("BM006", "BuildMat Brasov", "Brasov", 45.658, 25.601, "Depozit", 3200),
            ("BM007", "BuildMat Oradea", "Oradea", 47.047, 21.919, "Depozit mic", 2100),
            ("BM008", "BuildMat Craiova", "Craiova", 44.330, 23.795, "Depozit", 2500),
            ("BM009", "BuildMat B2B Portal", "Bucuresti", 44.437, 26.097, "Online B2B", 0),
            ("BM010", "BuildMat Ploiesti", "Ploiesti", 44.940, 26.022, "Depozit mic", 1800),
        ),
        produse=_prod(
            ("MT-01", "Ciment Portland 25kg", "Materiale", "Ciment", "CemRo", 28.0),
            ("MT-02", "Caramida plina", "Materiale", "Zidarie", "BrickPro", 2.5),
            ("MT-03", "Gresie 60x60 cm", "Finisaje", "Gresie", "TileArt", 89.0),
            ("MT-04", "Vopsea lavabila 15L", "Finisaje", "Vopsea", "ColorMax", 245.0),
            ("MT-05", "Profile rigips", "Materiale", "Gips-carton", "DryWall", 18.0),
            ("UN-01", "Ciocan cu maner", "Unelte", "Manual", "ToolPro", 45.0),
            ("UN-02", "Bormasina cu acumulator", "Unelte", "Electrice", "ToolPro", 399.0),
            ("UN-03", "Nivel laser", "Unelte", "Masura", "LaserLine", 289.0),
            ("MT-06", "Polistiren 10cm", "Izolatie", "Termic", "ThermoPlus", 42.0),
            ("MT-07", "Adeziv gresie 25kg", "Finisaje", "Adezivi", "FixTile", 38.0),
            ("MT-08", "Teava PVC 110", "Instalatii", "PVC", "PipeRo", 55.0),
            ("MT-09", "Tabla zincata", "Acoperis", "Metal", "RoofSteel", 120.0),
            ("UN-04", "Set chei combinate", "Unelte", "Manual", "ToolPro", 159.0),
            ("MT-10", "Betoniera 180L", "Utilaje mici", "Beton", "MixMaster", 1890.0),
            ("MT-11", "OSB 12mm placa", "Materiale", "Lemn", "WoodBoard", 72.0),
            ("MT-12", "Silicon constructii", "Finisaje", "Sigilare", "SealPro", 24.0),
            ("UN-05", "Echipament protectie set", "Unelte", "EPI", "SafeBuild", 189.0),
            ("MT-13", "Saci rumegus", "Materiale", "Agregate", "AgroFill", 15.0),
        ),
        clienti=_cli(
            ("BC001", "SC Construct Delta SRL", "PJ", "Bucuresti", "Constructor", 8.4),
            ("BC002", "SC Home Build SA", "PJ", "Cluj-Napoca", "Constructor", 8.7),
            ("BC003", "Ion Vasilescu", "PF", "Timisoara", "Bricolaj", 7.6),
            ("BC004", "SC Road Works SRL", "PJ", "Constanta", "Constructor", 8.1),
            ("BC005", "Maria Ionescu", "PF", "Brasov", "Bricolaj", 8.9),
            ("BC006", "SC Apartment Reno SRL", "PJ", "Iasi", "Renovari", 8.0),
            ("BC007", "SC Industrial Hall SRL", "PJ", "Oradea", "Constructor", 8.5),
            ("BC008", "Andrei Pop", "PF", "Sibiu", "Bricolaj", 7.1),
            ("BC009", "SC Bridge Team SRL", "PJ", "Craiova", "Constructor", 7.9),
            ("BC010", "Elena Radu", "PF", "Galati", "Bricolaj", 8.2),
            ("BC011", "SC Solar Install SRL", "PJ", "Ploiesti", "Instalator", 8.3),
            ("BC012", "SC Mega Developer SA", "PJ", "Bucuresti", "Constructor", 9.0),
            ("BC013", "Catalin Moldovan", "PF", "Braila", "Bricolaj", 6.8),
            ("BC014", "SC Pool Builder SRL", "PJ", "Pitesti", "Renovari", 8.1),
            ("BC015", "Diana Luca", "PF", "Voluntari", "Bricolaj", 8.6),
            ("BC016", "SC Timber Frame SRL", "PJ", "Targu Mures", "Constructor", 7.7),
            ("BC017", "Radu Enache", "PF", "Alba Iulia", "Bricolaj", 7.5),
            ("BC018", "SC Mining Infra SRL", "PJ", "Baia Mare", "Constructor", 8.0),
            ("BC019", "Simona Nita", "PF", "Pantelimon", "Bricolaj", 9.0),
            ("BC020", "SC Architect Studio SRL", "PJ", "Cluj-Napoca", "Proiectant", 8.2),
            ("BC021", "Victor Marinescu", "PF", "Oradea", "Bricolaj", 6.9),
            ("BC022", "SC Public Works SRL", "PJ", "Bucuresti", "Constructor", 8.8),
            ("BC023", "Ana Cojocaru", "PF", "Iasi", "Bricolaj", 7.4),
            ("BC024", "SC Harbor Build SRL", "PJ", "Constanta", "Constructor", 8.1),
            ("BC025", "George Ilie", "PF", "Craiova", "Bricolaj", 8.7),
            ("BC026", "SC Mountain Chalet SRL", "PJ", "Brasov", "Renovari", 8.3),
            ("BC027", "Laura Sandu", "PF", "Timisoara", "Bricolaj", 8.0),
            ("BC028", "SC Electric Install SRL", "PJ", "Sibiu", "Instalator", 8.4),
        ),
        canal=("Depozit", "Livrare santier", "Online B2B", "Telefon comanda"),
        online_mag_id=9,
        pages=(
            (
                "Performanță vânzări",
                "Pagina de deschidere pentru directorul comercial: total, timp, depozite.",
                (
                    "Care este valoarea totală a vânzărilor?",
                    "Cum evoluează vânzările pe luni (sezon construcții)?",
                    "Ce regiuni au depozite mai active?",
                    "Top depozite după valoare?",
                ),
                (
                    "Recomandare: KPI-uri + diagramă linie pe calendar.",
                    "Recomandare: coloane pe depozit sau regiune.",
                ),
            ),
            (
                "Mix produse",
                "Materiale grele, finisaje, unelte — ce mișcă stocul.",
                (
                    "Materiale vs Unelte vs Finisaje — cum stau la valoare?",
                    "Ce subcategorii (Ciment, Gresie, Manual…) sunt cele mai cerute?",
                    "Care SKU-uri sunt în top ca încasări?",
                ),
                (
                    "Recomandare: stacked column pe categorie.",
                    "Recomandare: bar chart produse; opțional filtru categorie.",
                ),
            ),
            (
                "Clienți",
                "Constructori, renovatori, bricolaj (PF).",
                (
                    "PJ vs PF — cine aduce bulk-ul vânzărilor?",
                    "Segmente (Constructor, Bricolaj, Renovari, Instalator): diferențe?",
                    "Care firme constructori sunt clienții cei mai mari?",
                ),
                (
                    "Recomandare: donut Tip_Client sau Segment.",
                    "Recomandare: top clienți PJ — tabel sau bar chart.",
                ),
            ),
            (
                "Depozite și livrări",
                "Ridicare din depozit vs livrare șantier vs portal B2B.",
                (
                    "Depozitele mari vs mici — aceeași dinamică?",
                    "Canal Depozit vs Livrare santier vs Online B2B — pondere?",
                    "Unde e puternică livrarea pe șantier (locații)?",
                ),
                (
                    "Recomandare: bar chart pe Nume_Magazin (depozit).",
                    "Recomandare: coloane pe Canal_Vanzare; matrice opțională depozit × canal.",
                ),
            ),
        ),
    ),
]


def build_fact_rows(team: TeamSpec, rng: random.Random, n: int = 128) -> list[list]:
    hdr = g5.FACT_HDR[:]
    title = [team.fact_title, *([""] * (len(hdr) - 1))]
    sub = [team.fact_sub, *([""] * (len(hdr) - 1))]
    max_days = (g5.CAL_END - g5.CAL_START).days
    rows: list[list] = []
    canal_primary = team.canal[0]
    for _ in range(n):
        d = g5.CAL_START + timedelta(days=rng.randint(0, max_days))
        id_client = rng.randint(1, len(team.clienti))
        id_produs = rng.randint(1, len(team.produse))
        id_magazin = rng.randint(1, len(team.magazine))
        id_data = g5.calendar_id(d)
        _, _, _, _, _, pret = team.produse[id_produs - 1]
        qty = rng.randint(1, 15)
        if rng.random() < 0.09:
            qty_str = f"{qty} buc"
        elif rng.random() < 0.05:
            qty_str = str(qty).replace(".", ",") + " buc"
        else:
            qty_str = qty
        disc = rng.choice([0, 0, 5, 10, 12, 15, 18])
        net = round(pret * qty * (1 - disc / 100), 2)
        tva_rate = 19 if rng.random() > 0.04 else 9
        tva = round(net * tva_rate / 100, 2)
        total = round(net + tva, 2)
        ds = g5.fmt_date_ro(d)
        if rng.random() < 0.04:
            ds = f" {ds} "
        id_client_out: str | int = id_client
        id_produs_out: str | int = id_produs
        id_magazin_out: str | int = id_magazin
        id_date_out: str | int = id_data
        if rng.random() < 0.05:
            id_client_out = f" {id_client} "
        if rng.random() < 0.03:
            id_produs_out = str(id_produs)
        if rng.random() < 0.03:
            id_magazin_out = f"{id_magazin}.0"
        pret_out = round(float(pret), 2)
        net_out = net
        total_out = total
        canal = rng.choice(team.canal)
        if id_magazin == team.online_mag_id and canal == canal_primary:
            canal = team.canal[1] if len(team.canal) > 1 else canal
        moneda = "RON" if rng.random() > 0.07 else "EUR"
        status = rng.choice(STATUS_MURDAR if rng.random() < 0.12 else STATUS)
        if status.strip().lower() == "anulat" and rng.random() < 0.5:
            net_out = 0
            total_out = 0
        rows.append(
            [
                id_date_out,
                id_client_out,
                id_produs_out,
                id_magazin_out,
                ds,
                qty_str,
                rng.choice(["buc", "BUC", "buc."]),
                pret_out,
                disc,
                net_out,
                tva_rate,
                tva,
                total_out,
                canal,
                status,
                moneda,
                "",
            ]
        )
    dup = rows[11][:]
    rows.insert(44, dup)
    rows[2] = ["" if c == "" else c for c in rows[2]]
    rows[2][1] = ""
    rows[6][4] = "31.02.2024"
    rows[8][4] = "2024-13-40"
    rows[10][4] = "15 mar 2023"
    rows[17][8] = "10%"
    rows[24][5] = "trei"
    rows[33][1] = f"  {rows[33][1]}  "
    rows[52] = [""] * len(hdr)
    rows[58] = ["SUBTOTAL ZILNIC", *([""] * (len(hdr) - 1))]
    rows[85][1] = 1
    rows[85][4] = rows[11][4]
    rows[98][5] = "3,5 buc"
    rows[108][13] = f"  {team.canal[-1]}  "
    rows[72][11] = "19,0"
    footer = [
        f"Total linii export: {len([r for r in rows if any(str(c).strip() for c in r)])}",
        *([""] * (len(hdr) - 1)),
    ]
    return [title, sub, hdr, *rows, footer]


def build_dim_clienti(team: TeamSpec, rng: random.Random) -> list[list]:
    hdr = [
        "ID",
        "Denumire",
        "Tip_Client",
        "Oras",
        "Judet",
        "Regiune",
        "Segment",
        "Data_Inregistrare",
        "NPS_Ultim",
        "Email_Contact",
    ]
    junk1 = [team.client_j1, *([""] * (len(hdr) - 1))]
    junk2 = [team.client_j2, *([""] * (len(hdr) - 1))]
    rows: list[list] = []
    for cid, (_cod, den, tip, oras, seg, nps) in enumerate(team.clienti, start=1):
        jud = g5.JUDET.get(oras, "N/A")
        reg = next(k for k, v in g5.ORASE.items() if oras in v)
        reg_out = reg if rng.random() > 0.08 else reg.upper()
        oras_out = oras + "\n" if rng.random() < 0.07 else oras
        tip_out = tip if rng.random() > 0.1 else ("Persoana Juridica" if tip == "PJ" else "Persoana Fizica")
        slug = den.split()[-1].lower().replace(".", "").replace("srl", "co")[:12]
        email = f"contact@{slug}.ro"
        if rng.random() < 0.06:
            email = f" {email} "
        d_reg = date(2018, 1, 1) + timedelta(days=rng.randint(0, 2500))
        d_str = d_reg.strftime("%d.%m.%Y") if rng.random() > 0.15 else d_reg.strftime("%Y-%m-%d")
        nps_out = nps if rng.random() > 0.05 else f"{nps}/10"
        id_out: str | int = cid
        if rng.random() < 0.04:
            id_out = str(cid)
        rows.append([id_out, den, tip_out, oras_out, jud, reg_out, seg, d_str, nps_out, email])
    dup_row = rows[4][:]
    dup_row[1] = team.dup_client_name
    dup_row[8] = "7,8"
    rows.append(dup_row)
    rows[4] = [rows[4][0], rows[4][1], rows[4][2], "CONSTANTA", rows[4][4], "Sud-Est", rows[4][6], rows[4][7], "-", rows[4][9]]
    rows[7] = [rows[7][0], rows[7][1], rows[7][2], rows[7][3], rows[7][4], rows[7][5], "Corporat", rows[7][7], rows[7][8], rows[7][9]]
    footer = ["— export terminat —", *([""] * (len(hdr) - 1))]
    return [junk1, junk2, hdr, *rows, footer]


def build_dim_produse(team: TeamSpec, rng: random.Random) -> list[list]:
    hdr = [
        "ID",
        "SKU_Produs",
        "Nume_Produs",
        "Categorie",
        "Subcategorie",
        "Brand",
        "Pret_Catalog",
        "Greutate_Kg",
        "Activ",
    ]
    rows: list[list] = []
    prev_cat = ""
    for pid, (sku, nume, cat, sub, brand, pret) in enumerate(team.produse, start=1):
        cat_cell = cat
        if rng.random() < 0.26:
            cat_cell = "" if prev_cat == cat else cat
        prev_cat = cat
        sku_out = "0" + sku if rng.random() < 0.04 else sku
        greutate = round(rng.uniform(0.05, 12.0), 2)
        activ = "Da" if rng.random() > 0.05 else "NU"
        nume_out = nume if rng.random() > 0.06 else f"  {nume}  "
        rows.append([pid, sku_out, nume_out, cat_cell, sub, brand, round(float(pret), 2), greutate, activ])
    if isinstance(rows[4][6], (int, float)):
        rows[4][6] = f"{rows[4][6]:.2f}".replace(".", ",") + " lei"
    rows[9][8] = 1
    rows[14] = [rows[14][0], rows[14][1], rows[14][2], rows[14][3], rows[14][4], rows[14][5], rows[14][6], rows[14][7], "YES"]
    header_row = ["ID", "Cod produs", "Denumire", "Cat.", "Subcat.", "Marca", "Pret", "Kg", "Activ?"]
    note = ["Categorii: fill-down manual daca lipseste", *([""] * (len(hdr) - 1))]
    return [[], header_row, hdr, *rows, note]


def build_dim_magazine(team: TeamSpec, rng: random.Random) -> list[list]:
    hdr = [
        "ID",
        "Nume_Magazin",
        "Oras",
        "Judet",
        "Regiune",
        "Latitudine",
        "Longitudine",
        "Format_Magazin",
        "Suprafata_mp",
    ]
    rows: list[list] = []
    reg_typos = {
        "Nord-Vest": "Nord Vest",
        "Centru": "Centru ",
        "Sud": "Sud",
        "Sud-Est": "Sud Est",
        "Bucuresti-Ilfov": "Bucuresti Ilfov",
    }
    for mid, (_cod, nume, oras, lat, lon, fmt, sup) in enumerate(team.magazine, start=1):
        jud = g5.JUDET.get(oras, "")
        reg = next(k for k, v in g5.ORASE.items() if oras in v)
        if rng.random() < 0.14:
            reg = reg_typos.get(reg, reg)
        lat_out: str | float = lat
        lon_out: str | float = lon
        if rng.random() < 0.2:
            lat_out = str(lat).replace(".", ",")
            lon_out = str(lon).replace(".", ",")
        sup_out: str | int = sup if rng.random() > 0.05 else f"{sup} mp"
        rows.append([mid, nume, oras, jud, reg, lat_out, lon_out, fmt, sup_out])
    rows[1][4] = "Bucuresti Ilfov"
    rows[0][8] = "N/A"
    title = [team.mag_sheet_title, *([""] * (len(hdr) - 1))]
    sub = ["Coordonate WGS84 — verificati semnul", *([""] * (len(hdr) - 1))]
    return [title, sub, hdr, *rows]


LUNI_EN = ["", "January", "February", "March", "April", "May", "June", "July", "August", "September", "October", "November", "December"]


def build_dim_date_team(rng: random.Random) -> list[list]:
    hdr = [
        "ID",
        "Data",
        "An",
        "Trimestru",
        "Luna",
        "Luna_Nume",
        "Zi_Saptamana",
        "Nr_Saptamana_An",
        "Zi_Lucratoare",
        "Perioada_Fiscala",
    ]
    rows: list[list] = []
    d = g5.CAL_START
    day_id = 1
    while d <= g5.CAL_END:
        q = (d.month - 1) // 3 + 1
        week = d.isocalendar()[1]
        zi = g5.ZI_RO[d.weekday()]
        lucr = "Da" if d.weekday() < 5 else "Nu"
        if rng.random() < 0.04:
            lucr = 1 if lucr == "Da" else 0
        pf = f"Q{q}-{d.year}" if rng.random() > 0.08 else f"T{q} {d.year}"
        if rng.random() < 0.02:
            data_out: str | int = (d - date(1899, 12, 30)).days
        elif rng.random() < 0.04:
            data_out = d.strftime("%d-%m-%Y")
        else:
            data_out = d.strftime("%d.%m.%Y")
        luna_nume = g5.LUNI_RO[d.month]
        if rng.random() < 0.03:
            luna_nume = LUNI_EN[d.month]
        trim_out = f"T{q}" if rng.random() > 0.1 else f"Trim {q}"
        rows.append([day_id, data_out, d.year, trim_out, d.month, luna_nume, zi, week, lucr, pf])
        day_id += 1
        d += timedelta(days=1)
    rows[120][1] = "32.01.2022"
    rows[600][5] = "Martie"
    rows[900][8] = "Y"
    title = ["Calendar vanzari · granularitate zi", *([""] * (len(hdr) - 1))]
    return [title, hdr, *rows]


def build_bundle(team: TeamSpec, rng: random.Random) -> dict[str, list[list]]:
    date_rng = random.Random(team.seed + 9000)
    return {
        "Vanzari": build_fact_rows(team, rng),
        "Dim_Clienti": build_dim_clienti(team, rng),
        "Dim_Produse": build_dim_produse(team, rng),
        "Dim_Magazine": build_dim_magazine(team, rng),
        "Dim_Date": build_dim_date_team(date_rng),
    }


def write_team_workbook(path: Path, bundle: dict[str, list[list]]) -> None:
    wb = Workbook()
    ws_fact = wb.active
    ws_fact.title = "Vanzari"
    g5.sheet_from_rows(ws_fact, bundle["Vanzari"])
    for sheet in ("Dim_Clienti", "Dim_Produse", "Dim_Magazine", "Dim_Date"):
        ws = wb.create_sheet(sheet)
        g5.sheet_from_rows(ws, bundle[sheet])
    wb.save(path)


def write_separate(team_dir: Path, bundle: dict[str, list[list]]) -> None:
    specs = [
        ("Vanzari_murdar.xlsx", "Vanzari"),
        ("Dim_Clienti_murdar.xlsx", "Dim_Clienti"),
        ("Dim_Produse_murdar.xlsx", "Dim_Produse"),
        ("Dim_Magazine_murdar.xlsx", "Dim_Magazine"),
        ("Dim_Date_murdar.xlsx", "Dim_Date"),
    ]
    for fname, sheet in specs:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet[:31]
        g5.sheet_from_rows(ws, bundle[sheet])
        wb.save(team_dir / fname)


def write_readme(team_dir: Path, team: TeamSpec) -> None:
    text = f"""Set date · {team.label} · Sesiunea 5
=====================================

Companie fictivă: {team.company} ({team.industry})
Context și cerințe raport: proiect-echipa.html

Fișier principal: Vanzari_s5_murdar.xlsx (5 foi)
Alternativ: fișiere separate *_murdar.xlsx

Structură: fact Vanzari + dimensiuni Dim_Clienti, Dim_Produse, Dim_Magazine, Dim_Date.
Set propriu echipei — probleme de calitate diferite față de exercițiul comun din folderul părinte.
"""
    (team_dir / "CITESTE_MAI_INTAI_s05_date.txt").write_text(text, encoding="utf-8")


def write_proiect_html(team_dir: Path, team: TeamSpec) -> None:
    """Brief-ul echipei; șablonul e în build_team_briefs.py."""
    from build_team_briefs import write_brief

    write_brief(team_dir, team)


def generate_team(team: TeamSpec) -> Path:
    team_dir = SESSION_S05 / team.folder
    team_dir.mkdir(parents=True, exist_ok=True)
    rng = random.Random(team.seed)
    bundle = build_bundle(team, rng)
    write_team_workbook(team_dir / "Vanzari_s5_murdar.xlsx", bundle)
    write_separate(team_dir, bundle)
    write_readme(team_dir, team)
    write_proiect_html(team_dir, team)
    return team_dir


def main() -> None:
    for team in TEAMS:
        path = generate_team(team)
        print(f"OK: {team.folder} -> {path}")


if __name__ == "__main__":
    main()
