# -*- coding: utf-8 -*-
"""Proiecte finale · Sesiunea 12: 6 teme pe companii reale, date livrate murdare în Excel.

Pentru fiecare temă citește assets/s12/ref/<tema>.json (datele găsite online: magazine, proiecte, branduri,
produse, prețuri, cifre de afaceri, surse) și generează:
  sesiuni/s12/proiecte/<tema>/date/*.xlsx      câte un fișier pe tabel, murdar (2 fapte + dimensiuni + Dim_Date)
  sesiuni/s12/proiecte/<tema>/CITESTE_MAI_INTAI.txt
  sesiuni/s12/proiecte/<tema>/index.html       pagina proiectului (vezi pagini_s12.py)
  sesiuni/s12/proiecte/index.html              pagina cu toate temele
  assets/s12/proiecte_solutii.md               cheia facilitatorului (rânduri curate, ce e murdar, relații)

Dimensiunile sunt reale; tranzacțiile sunt simulate, dar calibrate pe trendul cifrelor publicate.
Murdăria e cea din proiectele S6–S11. Dim_Date e curată și identică la toate temele.
Rulare: python generate_proiecte_s12.py
"""
from __future__ import annotations

import json
import math
import random
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
COURSE = ROOT.parent.parent
OUT = COURSE / "sesiuni" / "s12" / "proiecte"
KEY_OUT = ROOT / "proiecte_solutii.md"

D0, D1 = date(2023, 1, 1), date(2025, 12, 31)
LUNI = ["Ianuarie", "Februarie", "Martie", "Aprilie", "Mai", "Iunie", "Iulie", "August", "Septembrie",
        "Octombrie", "Noiembrie", "Decembrie"]
ZILE = ["Luni", "Marți", "Miercuri", "Joi", "Vineri", "Sâmbătă", "Duminică"]
PRENUME = ["Andrei", "Alexandru", "Mihai", "Ionuț", "Cristian", "Vlad", "Radu", "Bogdan", "Florin", "Adrian",
           "Ștefan", "Gabriel", "Daniel", "Marius", "Sorin", "Ana", "Maria", "Elena", "Ioana", "Andreea",
           "Alexandra", "Cristina", "Raluca", "Diana", "Roxana", "Mihaela", "Irina", "Laura", "Bianca", "Oana"]
NUME = ["Popescu", "Ionescu", "Popa", "Dumitru", "Stan", "Stoica", "Gheorghe", "Matei", "Ciobanu", "Rusu",
        "Munteanu", "Constantin", "Marin", "Dinu", "Tudor", "Lazăr", "Florea", "Barbu", "Neagu", "Moldovan",
        "Toma", "Sima", "Ilie", "Nistor", "Cozma", "Mocanu", "Preda", "Diaconu", "Enache", "Vasile"]


# ---------------------------------------------------------------------------
# Utilitare
# ---------------------------------------------------------------------------
def load_ref(tema: str) -> dict:
    return json.loads((ROOT / "ref" / f"{tema}.json").read_text(encoding="utf-8"))


def days(a: date, b: date) -> list[date]:
    return [a + timedelta(d) for d in range((b - a).days + 1)]


def months(a: date, b: date) -> list[date]:
    out, d = [], date(a.year, a.month, 1)
    while d <= b:
        out.append(d)
        d = date(d.year + (d.month == 12), d.month % 12 + 1, 1)
    return out


def pick(rng: random.Random, weights: dict):
    return rng.choices(list(weights), weights=list(weights.values()))[0]


def person(rng: random.Random) -> str:
    return f"{rng.choice(PRENUME)} {rng.choice(NUME)}"


def r2(x: float) -> float:
    return round(x, 2)


def rand_day(rng: random.Random, year: int, month_w: list[float], end: date = D1) -> date:
    while True:
        m = rng.choices(range(1, 13), weights=month_w)[0]
        dim = (date(year + (m == 12), m % 12 + 1, 1) - timedelta(1)).day
        d = date(year, m, rng.randint(1, dim))
        if D0 <= d <= end:
            return d


def fill_to_targets(rng, targets: dict, make_line, value_key: str) -> list[dict]:
    """Generează rânduri an cu an până când suma valorilor atinge ținta anului (păstrează prețurile reale)."""
    rows = []
    for year, target in targets.items():
        s = 0.0
        while s < target:
            r = make_line(year)
            if r is None:
                continue
            rows.append(r)
            s += r[value_key]
    return rows


def calibrate(rng, trend: dict, make_line, value_key: str, n_rows: int) -> tuple[list[dict], float]:
    """Alege ponderea din CA reală (eșantionul) astfel încât să iasă ~n_rows rânduri, apoi umple pe ani."""
    pilot = [make_line(y) for y in trend for _ in range(400)]
    avg = sum(r[value_key] for r in pilot if r) / sum(1 for r in pilot if r)
    share = n_rows * avg / sum(trend.values())
    rows = fill_to_targets(rng, {y: v * share for y, v in trend.items()}, make_line, value_key)
    return rows, share


def renumber(rows: list[dict], key: str) -> list[dict]:
    rows.sort(key=lambda r: [r[k] for k in r if isinstance(r[k], (date, datetime))][:1] or [0])
    for i, r in enumerate(rows, 1):
        r[key] = i
    return rows


def first_date(r: dict):
    for v in r.values():
        if isinstance(v, date):
            return v
    return date(2000, 1, 1)


def order(rows: list[dict], key: str) -> list[dict]:
    rows = sorted(rows, key=first_date)
    out = []
    for i, r in enumerate(rows, 1):
        out.append({key: i, **{k: v for k, v in r.items() if k != key}})
    return out


# ---------------------------------------------------------------------------
# Dim_Date: aceeași, curată, la toate temele
# ---------------------------------------------------------------------------
def gen_dim_date() -> list[dict]:
    rows = []
    for d in days(D0, D1):
        rows.append({
            "Data": d, "An": d.year, "Trimestru": f"T{(d.month - 1) // 3 + 1}", "Luna": d.month,
            "Luna_Nume": LUNI[d.month - 1], "Luna_Scurta": LUNI[d.month - 1][:3], "An_Luna": f"{d.year}-{d.month:02d}",
            "Saptamana_ISO": d.isocalendar()[1], "Zi_Saptamana": d.isoweekday(), "Zi_Nume": ZILE[d.weekday()],
            "Este_Weekend": "Da" if d.weekday() >= 5 else "Nu",
        })
    return rows


# ---------------------------------------------------------------------------
# Murdărire generică (aceleași tipuri ca în S6–S11)
# ---------------------------------------------------------------------------
def case_variant(rng, s, p):
    if not isinstance(s, str):
        return s
    x = rng.random()
    if x < p * 0.4:
        return s.lower()
    if x < p * 0.7:
        return s.upper()
    if x < p:
        return s + "  " if rng.random() < 0.5 else "  " + s
    return s


def fmt_num(v) -> str:
    if isinstance(v, float) and not v.is_integer():
        return f"{v:.2f}".replace(".", ",")
    return str(int(v))


def dirty(rows: list[dict], spec: dict, seed: int) -> tuple[list[list], dict]:
    """Întoarce (rândurile murdare cu antet, rezumatul murdăriei pentru cheie)."""
    rng = random.Random(seed)
    cols = list(rows[0])
    info = []
    out = []
    for r in rows:
        v = dict(r)
        for c in spec.get("case", []):
            v[c] = case_variant(rng, v[c], spec.get("case_p", 0.35))
        for c in spec.get("trim", []):
            if isinstance(v[c], str) and rng.random() < 0.3:
                v[c] = " " + v[c] + "  "
        for c, unit in spec.get("units", {}).items():
            if v[c] is not None:
                v[c] = f"{fmt_num(v[c])} {unit}"
        for c in spec.get("comma", []):
            if isinstance(v[c], (int, float)):
                v[c] = f"{v[c]:.2f}".replace(".", ",")
        for c in spec.get("pct", []):
            if isinstance(v[c], (int, float)):
                v[c] = f"{round(v[c] * 100)}%"
        for c in spec.get("date_text", []):
            if isinstance(v[c], date):
                v[c] = v[c].strftime("%d.%m.%Y")
        for c in spec.get("id_text", []):
            if v[c] is not None and rng.random() < 0.25:
                v[c] = f" {v[c]} "
        for c, default in spec.get("null_default", {}).items():
            if v[c] == default and rng.random() < 0.6:
                v[c] = None
        for c, mapping in spec.get("synonyms", {}).items():
            if v[c] in mapping and rng.random() < 0.3:
                v[c] = mapping[v[c]]
        out.append(v)
    for k, txt in (("case", "litere mari/mici amestecate și spații"), ("trim", "spații la început/sfârșit"),
                   ("units", "numere cu unități ca text"), ("comma", "zecimale cu virgulă (text)"),
                   ("pct", "procente ca text („15%”)"), ("date_text", "date ca text dd.mm.yyyy"),
                   ("id_text", "chei ca text cu spații"), ("null_default", "valori lipsă în loc de valoarea implicită"),
                   ("synonyms", "aceeași valoare scrisă în mai multe feluri")):
        if spec.get(k):
            names = list(spec[k]) if isinstance(spec[k], dict) else spec[k]
            info.append(f"{txt}: {', '.join(names)}")
    extra = spec.get("extra")
    if extra:
        for v in out:
            v[extra] = None if extra == "Observatii" else "2026-01-05 08:00"
        cols.append(extra)
        info.append(f"coloană în plus: {extra}")
    n_d = spec.get("dupes", 0)
    for i in sorted(rng.sample(range(3, len(out) - 3), n_d), reverse=True):
        out.insert(i + 1, dict(out[i]))
    if n_d:
        info.append(f"{n_d} duplicate")
    key = cols[0]
    n_t = spec.get("tests", 0)
    text_col = next((c for c in cols[1:] if isinstance(rows[0].get(c), str)), cols[1])
    for t in range(n_t):
        v = dict(out[t + 1])
        v[key] = f"TEST{t + 1:02d}"
        v[text_col] = "TEST - nu folositi"
        out.insert(rng.randint(len(out) // 3, len(out) - 2), v)
    if n_t:
        info.append(f"{n_t} rânduri de test ({key} = TEST…)")
    n_b = spec.get("blanks", 0)
    for _ in range(n_b):
        out.insert(rng.randint(5, len(out) - 5), {c: None for c in cols})
    if n_b:
        info.append(f"{n_b} rânduri goale")
    if spec.get("total"):
        tot = {c: None for c in cols}
        tot[key] = "TOTAL"
        for c in spec["total"]:
            tot[c] = r2(sum(r[c] for r in rows if isinstance(r[c], (int, float))))
        out.append(tot)
        info.append("rând TOTAL la final")
    if spec.get("title"):
        info.append("rând de titlu deasupra antetului")
    body = [[v.get(c) for c in cols] for v in out]
    header = [spec["title"]] if spec.get("title") else None
    return ([header] if header else []) + [cols] + body, {"info": info, "cols_dirty": len(cols)}


def check_table(clean: list[dict], dirty_rows: list[list], spec: dict) -> None:
    """Curățarea minimă (fără titlu, rânduri goale, test, TOTAL, duplicate) trebuie să dea exact rândurile curate."""
    body = dirty_rows[2:] if spec.get("title") else dirty_rows[1:]
    seen, ids = set(), []
    for row in body:
        if all(c is None for c in row):
            continue
        k = str(row[0]).strip()
        if k == "TOTAL" or k.startswith("TEST"):
            continue
        sig = tuple(str(c).strip().lower() if c is not None else None for c in row)
        if sig in seen:
            continue
        seen.add(sig)
        ids.append(int(k))
    want = [r[list(r)[0]] for r in clean]
    assert sorted(ids) == sorted(want), f"check eșuat: {len(ids)} vs {len(want)}"


# ---------------------------------------------------------------------------
# Scriere Excel
# ---------------------------------------------------------------------------
def write_xlsx(path: Path, sheet: str, rows: list[list]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]
    for r in rows:
        ws.append(r)
    width = max(len(r) for r in rows)
    for i in range(1, width + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, datetime) or isinstance(c.value, date):
                c.number_format = "yyyy-mm-dd"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


# ===========================================================================
# 1. TELEKOM ROMANIA BUSINESS
# ===========================================================================
def build_telekom() -> dict:
    ref = load_ref("telekom")
    rng = random.Random(1201)
    end = date(2025, 9, 30)
    eur = ref["curs_bnr_mediu"]["EUR"]
    reg_of = {j: r for r, js in ref["regiuni_vanzari"].items() for j in js}
    w_judet = {j: (14 if j == "București" else 2.5 if j in ("Ilfov", "Cluj", "Timiș", "Iași", "Constanța", "Brașov", "Prahova") else 1)
               for j in reg_of}

    seg_w = {"SOHO": 0.38, "SME": 0.42, "Corporate": 0.12, "Sector public": 0.08}
    clienti = []
    for i in range(1, 651):
        seg = pick(rng, seg_w)
        ind = "Administrație publică" if seg == "Sector public" else rng.choice([x for x in ref["industrii_clienti"] if x != "Administrație publică"])
        j = pick(rng, w_judet)
        ang = {"SOHO": "1-9", "SME": rng.choice(["10-49", "50-249"]), "Corporate": rng.choice(["250-999", "1000+"]), "Sector public": rng.choice(["50-249", "250-999"])}[seg]
        clienti.append({"ID_Client": i, "Cod_Client": f"CLB-{i:04d}", "Segment": seg, "Industrie": ind, "Judet": j,
                        "Regiune": reg_of[j], "Nr_Angajati": ang,
                        "Client_Din": date(rng.randint(2012, 2024), rng.randint(1, 12), 1)})

    produse = []
    for a in ref["abonamente_business"]:
        produse.append({"Produs": a["nume"], "Categorie": "Abonament mobil", "Vendor": "Telekom", "Tip_Venit": "Recurent",
                        "Pret_Lista_EUR": a["pret_eur_luna_fara_tva"], "Unitate": "SIM / lună"})
    for s in ref["solutii_ict"]:
        pret = next(v for k, v in s.items() if k.startswith("pret_"))
        produse.append({"Produs": s["nume"], "Categorie": s["tip"], "Vendor": s["vendor"], "Tip_Venit": "Recurent",
                        "Pret_Lista_EUR": pret, "Unitate": "licență / lună" if "Licen" in s["tip"] else "dispozitiv / lună"})
    for e in ref["echipamente"]:
        produse.append({"Produs": e["nume"], "Categorie": "Echipament " + e["categorie"].lower(), "Vendor": e["vendor"],
                        "Tip_Venit": "One-off", "Pret_Lista_EUR": e["pret_eur_fara_tva"], "Unitate": "bucată"})
    produse = [{"ID_Produs": i, **p} for i, p in enumerate(produse, 1)]

    servicii = []
    ore = {"Instalare": (2, 6), "Deployment": (4, 16), "Integrare": (8, 40), "Configurare": (4, 12),
           "Montaj": (1, 3), "Migrare": (8, 60), "Mentenanță": (1, 4), "Audit": (16, 80)}
    for i, s in enumerate(ref["servicii_tehnice"], 1):
        cuv = s["nume"].split()[0]
        servicii.append({"ID_Serviciu": i, "Serviciu": s["nume"], "Tip_Serviciu": cuv,
                         "Tarif_EUR_Ora": s["tarif_eur_ora"], "SLA_Ore": 72 if cuv in ("Integrare", "Audit", "Migrare") else 48 if cuv != "Mentenanță" else 24})

    echipa = []
    for i in range(1, 71):
        if i <= 46:
            rol = pick(rng, {"Key Account Manager": 0.3, "Account Manager SME": 0.5, "Telesales": 0.2})
        else:
            rol = pick(rng, {"Tehnician teren": 0.55, "Inginer rețele": 0.3, "Specialist securitate": 0.15})
        reg = pick(rng, {r: (4 if r == "București-Ilfov" else 1) for r in ref["regiuni_vanzari"]})
        echipa.append({"ID_Angajat": i, "Nume": person(rng), "Rol": rol, "Departament": "Vânzări B2B" if i <= 46 else "Suport tehnic",
                       "Regiune": reg, "Data_Angajare": date(rng.randint(2010, 2023), rng.randint(1, 12), rng.randint(1, 28))})
    agenti = [e for e in echipa if e["Departament"] == "Vânzări B2B"]
    tehnicieni = [e for e in echipa if e["Departament"] == "Suport tehnic"]

    qty = {"SOHO": (1, 5), "SME": (3, 25), "Corporate": (20, 250), "Sector public": (10, 120)}
    cat_w = {"Abonament mobil": 0.45, "Echipament": 0.32, "ICT": 0.23}
    month_w = [0.8, 0.85, 1, 1, 1, 0.95, 0.85, 0.7, 1.05, 1.1, 1.15, 1.3]

    def line(year):
        c = rng.choice(clienti)
        d = rand_day(rng, year, month_w, end)
        if d < c["Client_Din"]:
            return None
        grup = pick(rng, cat_w)
        pool = [p for p in produse if (p["Categorie"] == "Abonament mobil") == (grup == "Abonament mobil")
                and (p["Tip_Venit"] == "One-off") == (grup == "Echipament")]
        p = rng.choice(pool)
        lo, hi = qty[c["Segment"]]
        q = rng.randint(lo, hi)
        if p["Tip_Venit"] == "One-off" and p["Pret_Lista_EUR"] > 400:
            q = max(1, q // 3)
        pret = r2(p["Pret_Lista_EUR"] * eur[str(year)])
        disc = rng.choice([0, 0, 0.05, 0.1, 0.15, 0.2]) if c["Segment"] in ("Corporate", "Sector public") else rng.choice([0, 0, 0, 0.05, 0.1])
        durata = 24 if p["Categorie"] in ("Abonament mobil", "IoT/M2M", "Conectivitate") else 12 if p["Tip_Venit"] == "Recurent" else 0
        val = pret * q * (durata or 1) * (1 - disc)
        regs = [a for a in agenti if a["Regiune"] == c["Regiune"]] or agenti
        ag = rng.choice([a for a in regs if (a["Rol"] == "Key Account Manager") == (c["Segment"] in ("Corporate", "Sector public"))] or regs)
        canal = "KAM" if ag["Rol"] == "Key Account Manager" else "Telesales" if ag["Rol"] == "Telesales" else rng.choice(["Magazin business", "Account manager", "Online"])
        return {"ID_Vanzare": 0, "Data_Comanda": d, "ID_Client": c["ID_Client"], "ID_Produs": p["ID_Produs"],
                "ID_Agent": ag["ID_Angajat"], "Tip_Contract": pick(rng, {"Nou": 0.45, "Reînnoire": 0.4, "Upgrade": 0.15}),
                "Canal": canal, "Cantitate": q, "Pret_Unitar_RON": pret, "Durata_Contract_Luni": durata,
                "Discount": disc, "Valoare_Contract_RON": r2(val)}

    trend = {y: v * (0.75 if y == 2025 else 1) for y, v in
             {int(k): v for k, v in ref["calibrare"]["venituri_totale_ron"].items()}.items()}
    vanzari, share = calibrate(rng, trend, line, "Valoare_Contract_RON", 11500)
    vanzari = order(vanzari, "ID_Vanzare")

    serv_for = {"Echipament rețea": [3, 4], "Echipament securitate": [4], "Echipament iot": [5], "IoT/M2M": [5],
                "Software/Licență": [6], "Echipament router": [1], "Conectivitate": [1], "Abonament mobil": [2]}
    prob = {"Echipament rețea": 0.8, "Echipament securitate": 0.85, "Echipament iot": 0.7, "IoT/M2M": 0.5,
            "Software/Licență": 0.18, "Echipament router": 0.35, "Conectivitate": 0.3, "Abonament mobil": 0.06}
    lucrari = []
    pmap = {p["ID_Produs"]: p for p in produse}
    cmap = {c["ID_Client"]: c for c in clienti}
    for v in vanzari:
        cat = pmap[v["ID_Produs"]]["Categorie"]
        if cat not in prob or rng.random() > prob[cat]:
            continue
        sid = rng.choice(serv_for[cat])
        lucrari.append((v, sid))
    for v in rng.sample(vanzari, 700):
        lucrari.append((v, rng.choice([7, 7, 7, 8])))
    fact_s = []
    for v, sid in lucrari:
        s = servicii[sid - 1]
        c = cmap[v["ID_Client"]]
        dp = v["Data_Comanda"] + timedelta(rng.randint(2, 25) if sid != 7 else rng.randint(30, 300))
        if dp > end:
            continue
        lo, hi = ore[s["Tip_Serviciu"]]
        est = rng.randint(lo, hi) * (2 if c["Segment"] == "Corporate" and sid in (3, 4, 6, 8) else 1)
        if sid == 5:
            est = max(1, round(v["Cantitate"] * rng.uniform(0.4, 0.8)))
        status = pick(rng, {"Finalizată": 0.92, "Replanificată": 0.05, "Anulată": 0.03})
        real = 0 if status == "Anulată" else r2(est * math.exp(rng.gauss(0.07, 0.22)))
        loc = [t for t in tehnicieni if t["Regiune"] == c["Regiune"]] or tehnicieni
        t = rng.choice([x for x in loc if (x["Rol"] == "Specialist securitate") == (sid in (4, 8))] or loc)
        ore_interv = r2(rng.lognormvariate(math.log(s["SLA_Ore"] * 0.55), 0.45))
        tarif = r2(s["Tarif_EUR_Ora"] * eur[str(dp.year)])
        fact_s.append({"ID_Lucrare": 0, "Data_Programare": dp,
                       "Data_Finalizare": None if status == "Anulată" else dp + timedelta(math.ceil(real / 8) - 1 if real else 0),
                       "ID_Client": c["ID_Client"], "ID_Serviciu": sid, "ID_Tehnician": t["ID_Angajat"],
                       "Status": status, "Ore_Estimate": est, "Ore_Realizate": real, "Tarif_Ora_RON": tarif,
                       "Cost_RON": r2(real * tarif), "Ore_Pana_La_Interventie": ore_interv,
                       "SLA_Respectat": "Da" if ore_interv <= s["SLA_Ore"] else "Nu",
                       "Rating_Client": None if status != "Finalizată" else rng.choices([5, 4, 3, 2, 1], [0.48, 0.32, 0.12, 0.05, 0.03])[0]})
    fact_s = order(fact_s, "ID_Lucrare")

    tables = {
        "Fact_Vanzari": (vanzari, {"date_text": ["Data_Comanda"], "comma": ["Pret_Unitar_RON"], "pct": ["Discount"],
                                   "units": {"Cantitate": "buc"}, "case": ["Canal", "Tip_Contract"], "id_text": ["ID_Client"],
                                   "extra": "Exportat_La", "dupes": 7, "tests": 2, "blanks": 3, "total": ["Valoare_Contract_RON"]},
                         "fiecare linie de contract B2B: abonamente, licențe și echipamente vândute, cu agent și discount"),
        "Fact_Servicii": (fact_s, {"units": {"Ore_Estimate": "ore"}, "comma": ["Ore_Realizate", "Ore_Pana_La_Interventie"],
                                   "case": ["Status", "SLA_Respectat"], "date_text": ["Data_Programare"], "dupes": 5, "blanks": 2,
                                   "title": "Export lucrări tehnice - Telekom Business"},
                          "lucrările tehnice: instalare, deployment, integrare rețele, securitate, mentenanță (ore, SLA, rating)"),
        "Dim_Clienti": (clienti, {"case": ["Industrie", "Segment"], "trim": ["Judet"], "dupes": 4, "tests": 1,
                                  "synonyms": {"Judet": {"București": "Bucuresti", "Timiș": "Timis", "Iași": "Iasi"}}},
                        "clienții business: segment, industrie, județ, regiune"),
        "Dim_Produse": (produse, {"case": ["Vendor"], "units": {"Pret_Lista_EUR": "EUR"}, "dupes": 1, "tests": 1},
                        "abonamentele Business, soluțiile ICT/IoT și echipamentele, cu prețul de listă"),
        "Dim_Servicii": (servicii, {"units": {"Tarif_EUR_Ora": "EUR/h", "SLA_Ore": "h"}, "trim": ["Serviciu"]},
                         "serviciile tehnice și tariful orar, cu SLA-ul în ore"),
        "Dim_Echipa": (echipa, {"case": ["Rol"], "trim": ["Nume"], "date_text": ["Data_Angajare"], "dupes": 2},
                       "agenții de vânzări B2B și tehnicienii"),
    }
    return {"tema": "telekom", "share": share, "tables": tables, "dim_date_desc": "calendarul 2023–2025 (datele se opresc la 30.09.2025)",
            "relatii": [("Fact_Vanzari", "ID_Client", "Dim_Clienti"), ("Fact_Vanzari", "ID_Produs", "Dim_Produse"),
                        ("Fact_Vanzari", "ID_Agent", "Dim_Echipa"), ("Fact_Vanzari", "Data_Comanda", "Dim_Date"),
                        ("Fact_Servicii", "ID_Client", "Dim_Clienti"), ("Fact_Servicii", "ID_Serviciu", "Dim_Servicii"),
                        ("Fact_Servicii", "ID_Tehnician", "Dim_Echipa"), ("Fact_Servicii", "Data_Programare", "Dim_Date")]}


# ===========================================================================
# 2. CONSTRUCȚII ERBAȘU
# ===========================================================================
ERBASU_PERIOADE = {  # (start, sfârșit planificat) estimate după stadiul publicat al lucrărilor
    1: ("2022-06", "2027-06"), 2: ("2025-01", "2026-12"), 3: ("2023-06", "2026-12"), 4: ("2024-11", "2027-06"),
    5: ("2024-11", "2027-06"), 6: ("2024-06", "2027-09"), 7: ("2023-03", "2025-12"), 8: ("2023-01", "2026-12"),
    9: ("2023-09", "2026-06"), 10: ("2023-05", "2025-10"), 11: ("2023-10", "2026-03"), 12: ("2023-01", "2025-12"),
    13: ("2023-04", "2025-09"), 14: ("2023-02", "2025-06"), 15: ("2023-01", "2025-08"), 16: ("2023-06", "2026-06"),
    17: ("2025-06", "2028-06"), 18: ("2024-03", "2026-06"), 19: ("2026-02", "2030-12"),
}
FAZE = {  # (început, sfârșit) ca fracție din durata proiectului
    "Organizare de șantier": (0.0, 1.0), "Terasamente și fundații": (0.0, 0.3), "Structură beton armat": (0.1, 0.6),
    "Structură metalică": (0.3, 0.7), "Arhitectură și finisaje": (0.55, 1.0), "Instalații electrice": (0.4, 0.95),
    "Instalații sanitare": (0.4, 0.9), "Instalații HVAC": (0.45, 0.95), "Drumuri și sistematizare": (0.0, 1.0),
    "Instalații speciale (gaze medicale, curenți slabi)": (0.6, 1.0),
}
CAT_TIP = {  # ce categorii intră la fiecare tip de proiect
    "Infrastructură rutieră": ["Organizare de șantier", "Terasamente și fundații", "Drumuri și sistematizare", "Structură beton armat", "Instalații electrice"],
    "Infrastructură urbană": ["Organizare de șantier", "Terasamente și fundații", "Structură beton armat", "Drumuri și sistematizare", "Instalații electrice", "Instalații sanitare"],
    "Infrastructură metrou": ["Organizare de șantier", "Terasamente și fundații", "Structură beton armat", "Instalații electrice"],
}
ART_CAT = {
    "Organizare de șantier": ["Agregate 0-31,5", "Placă gips-carton 12,5 mm"],
    "Terasamente și fundații": ["Agregate 0-31,5", "Beton C25/30", "Oțel beton PC52 Ø14", "Țeavă PVC canalizare Ø110"],
    "Structură beton armat": ["Beton C30/37", "Beton C25/30", "Oțel beton PC52 Ø14", "Plasă sudată Ø6 100x100", "Ciment CEM II/A-S 42,5R"],
    "Structură metalică": ["Subcontractare structuri metalice"],
    "Arhitectură și finisaje": ["BCA 25 cm", "Bloc ceramic Porotherm 25", "Placă gips-carton 12,5 mm", "Polistiren EPS 80 10 cm", "Vată minerală bazaltică 10 cm", "Tâmplărie PVC/aluminiu", "Gresie porțelanată 60x60", "Vopsea lavabilă interior"],
    "Instalații electrice": ["Cablu CYY-F 3x2,5", "Cablu CYABY 4x16", "Tablou electric echipat", "Corp iluminat LED 40W", "Subcontractare instalații electrice"],
    "Instalații sanitare": ["Țeavă PPR Ø25", "Țeavă PVC canalizare Ø110", "Subcontractare instalații sanitare"],
    "Instalații HVAC": ["Centrală de tratare aer (CTA)", "Chiller 300 kW", "Tubulatură ventilație tablă zincată", "Subcontractare instalații HVAC"],
    "Drumuri și sistematizare": ["Mixtură asfaltică BA16", "Agregate 0-31,5", "Beton C25/30"],
    "Instalații speciale (gaze medicale, curenți slabi)": ["Subcontractare instalații speciale"],
}


def build_erbasu() -> dict:
    ref = load_ref("erbasu")
    rng = random.Random(1202)
    idx = ref["indice_cost_constructii_ins"]

    benef = {}
    for p in ref["proiecte"]:
        b = p["beneficiar"]
        if b not in benef:
            tip = "Privat" if "privat" in b.lower() else "Public central" if b.startswith(("Ministerul", "CNAIR", "CNI", "Academia")) else "Public local"
            benef[b] = {"ID_Beneficiar": len(benef) + 1, "Beneficiar": b, "Tip_Beneficiar": tip,
                        "Judet": p["judet"] if tip != "Public central" else "București"}
    beneficiari = list(benef.values())

    proiecte = []
    for i, p in enumerate(ref["proiecte"], 1):
        s, e = ERBASU_PERIOADE[i]
        proiecte.append({"ID_Proiect": i, "Proiect": p["nume"], "Tip_Proiect": p["tip"], "Oras": p["oras"], "Judet": p["judet"],
                         "ID_Beneficiar": benef[p["beneficiar"]]["ID_Beneficiar"], "Valoare_Contract_RON": p["valoare_ron"],
                         "Data_Start": date(int(s[:4]), int(s[5:]), 1), "Data_Sfarsit_Planificata": date(int(e[:4]), int(e[5:]), 28),
                         "Status": p["status"]})

    categorii = [{"ID_Categorie": i, "Categorie_Lucrare": c["nume"], "Grupa": "Instalații" if c["nume"].startswith("Instalații") else
                  "Structură" if c["nume"].startswith(("Structură", "Terasamente")) else "Infrastructură" if c["nume"].startswith("Drumuri") else "General",
                  "Pondere_Tipica": c["pondere"]} for i, c in enumerate(ref["categorii_lucrari"], 1)]
    cat_id = {c["Categorie_Lucrare"]: c["ID_Categorie"] for c in categorii}

    furnizori, prod_f = [], {}
    for m in ref["materiale"]:
        pr = m["producator"]
        if pr != "Producător local" and pr not in prod_f:
            prod_f[pr] = len(furnizori) + 1
            furnizori.append({"ID_Furnizor": len(furnizori) + 1, "Furnizor": pr, "Tip_Furnizor": "Producător materiale",
                              "Specializare": m["nume"].split()[0], "Judet": "București"})
    for j in ["Bihor", "Dolj", "Timiș", "Constanța", "București", "Argeș", "Tulcea", "Cluj", "Iași", "Prahova"]:
        furnizori.append({"ID_Furnizor": len(furnizori) + 1, "Furnizor": f"Furnizor local materiale {j}", "Tip_Furnizor": "Distribuitor local",
                          "Specializare": "Agregate, asfalt, tâmplărie", "Judet": j})
    sub_spec = {"structuri metalice": 2, "instalații electrice": 3, "instalații sanitare": 2, "instalații HVAC": 3, "instalații speciale": 2}
    for sp, n in sub_spec.items():
        for k in range(1, n + 1):
            furnizori.append({"ID_Furnizor": len(furnizori) + 1, "Furnizor": f"Subcontractor {sp} #{k}", "Tip_Furnizor": "Subcontractor",
                              "Specializare": sp, "Judet": rng.choice(["București", "Cluj", "Timiș", "Bihor", "Dolj"])})

    articole = [{"ID_Articol": i, "Articol": m["nume"], "Tip_Articol": "Material", "UM": m["um"], "Producator": m["producator"],
                 "Pret_Referinta_RON": m["pret_ron"]} for i, m in enumerate(ref["materiale"], 1)]
    for sp in sub_spec:
        articole.append({"ID_Articol": len(articole) + 1, "Articol": f"Subcontractare {sp}", "Tip_Articol": "Serviciu subcontractat",
                         "UM": "situație lucrări", "Producator": None, "Pret_Referinta_RON": None})
    art_by_name = {a["Articol"]: a for a in articole}

    situatii, achizitii = [], []
    for p in proiecte:
        s, e = p["Data_Start"], p["Data_Sfarsit_Planificata"]
        tot_m = len(months(s, e))
        cats = CAT_TIP.get(p["Tip_Proiect"], list(FAZE))
        w = {c: next(x["Pondere_Tipica"] for x in categorii if x["Categorie_Lucrare"] == c) for c in cats}
        sw = sum(w.values())
        lag = rng.uniform(0.8, 1.02)  # proiectele publice tind să rămână în urmă față de plan
        cum_plan = {c: 0.0 for c in cats}
        cum_exec = {c: 0.0 for c in cats}
        for k, m in enumerate(months(s, e)):
            if m > D1:
                break
            t0, t1 = k / tot_m, (k + 1) / tot_m
            nm = 0
            for c in cats:
                a, b = FAZE[c]
                if t1 <= a or t0 >= b:
                    continue
                # curbă în S pe faza categoriei
                f = lambda t: 0 if t <= a else 1 if t >= b else (1 - math.cos(math.pi * (t - a) / (b - a))) / 2
                plan = p["Valoare_Contract_RON"] * w[c] / sw * (f(t1) - f(t0))
                if plan < 1000:
                    continue
                exe = plan * lag * rng.uniform(0.7, 1.25)
                if m.month in (1, 2) and rng.random() < 0.5:
                    exe *= rng.uniform(0.3, 0.7)  # iarna se lucrează mai puțin
                cum_plan[c] += plan
                cum_exec[c] += exe
                if m < D0:
                    continue  # lunile dinainte de 2023 contează doar la avansul cumulat
                cat_val = p["Valoare_Contract_RON"] * w[c] / sw
                delay = max(0, round((cum_plan[c] - cum_exec[c]) / plan * 30)) if plan else 0
                situatii.append({"ID_Situatie": 0, "Luna": m, "ID_Proiect": p["ID_Proiect"], "ID_Categorie": cat_id[c],
                                 "Valoare_Planificata_RON": r2(plan), "Valoare_Executata_RON": r2(exe),
                                 "Avans_Cumulat": round(min(cum_exec[c] / cat_val, 1), 4), "Zile_Intarziere": min(delay, 120),
                                 "Nr_Muncitori": max(4, round(exe / rng.uniform(55000, 80000)))})
                nm += 1
                # achizițiile lunii pentru categoria asta: ~55% din execuție
                buget = exe * rng.uniform(0.45, 0.62)
                names = ART_CAT[c]
                n_lines = rng.randint(1, min(4, len(names) + 1))
                for nm_art in rng.sample(names, min(n_lines, len(names))):
                    a_ = art_by_name[nm_art]
                    part = buget / n_lines
                    dt = m + timedelta(rng.randint(0, 27))
                    if a_["Tip_Articol"] == "Material":
                        pret = r2(a_["Pret_Referinta_RON"] * idx[str(m.year)] * rng.uniform(0.93, 1.08))
                        q = max(1, round(part / pret))
                        fid = prod_f.get(a_["Producator"]) or rng.choice([f["ID_Furnizor"] for f in furnizori if f["Tip_Furnizor"] == "Distribuitor local"])
                        val = q * pret
                    else:
                        sp = nm_art.replace("Subcontractare ", "")
                        pret = r2(part)
                        q = 1
                        fid = rng.choice([f["ID_Furnizor"] for f in furnizori if f["Specializare"] == sp])
                        val = pret
                    achizitii.append({"ID_Achizitie": 0, "Data_Achizitie": dt, "ID_Proiect": p["ID_Proiect"],
                                      "ID_Articol": a_["ID_Articol"], "ID_Furnizor": fid, "Cantitate": q,
                                      "Pret_Unitar_RON": pret, "Valoare_RON": r2(val),
                                      "Termen_Plata_Zile": rng.choice([30, 30, 60, 60, 90])})
    situatii = order(situatii, "ID_Situatie")
    achizitii = order(achizitii, "ID_Achizitie")
    exec_an = {y: sum(s["Valoare_Executata_RON"] for s in situatii if s["Luna"].year == y) for y in (2023, 2024, 2025)}
    ca = {int(k): v for k, v in ref["calibrare"]["cifra_afaceri_ron"].items()}
    share = sum(exec_an.values()) / sum(ca.values())

    tables = {
        "Fact_Situatii_Lucrari": (situatii, {"comma": ["Valoare_Planificata_RON", "Valoare_Executata_RON"], "pct": ["Avans_Cumulat"],
                                             "units": {"Zile_Intarziere": "zile"}, "date_text": ["Luna"], "id_text": ["ID_Proiect"],
                                             "extra": "Observatii", "dupes": 6, "blanks": 3, "total": ["Valoare_Executata_RON"]},
                                  "situațiile lunare de lucrări pe proiect și categorie: planificat vs. executat, avans, întârziere"),
        "Fact_Achizitii": (achizitii, {"units": {"Cantitate": "buc"}, "comma": ["Pret_Unitar_RON"], "date_text": ["Data_Achizitie"],
                                       "id_text": ["ID_Furnizor"], "dupes": 8, "tests": 2, "blanks": 2, "extra": "Exportat_La"},
                           "achizițiile de materiale și serviciile subcontractate pe proiect"),
        "Dim_Proiecte": (proiecte, {"trim": ["Proiect"], "case": ["Status", "Tip_Proiect"], "date_text": ["Data_Start"], "dupes": 1},
                         "proiectele din portofoliu: tip, oraș, beneficiar, valoare contract, termene"),
        "Dim_Beneficiari": (beneficiari, {"case": ["Tip_Beneficiar"], "dupes": 1},
                            "autoritățile și investitorii pentru care se construiește"),
        "Dim_Categorii_Lucrari": (categorii, {"pct": ["Pondere_Tipica"], "trim": ["Categorie_Lucrare"]},
                                  "categoriile de lucrări: structură, instalații electrice, sanitare, HVAC, finisaje…"),
        "Dim_Articole": (articole, {"case": ["Tip_Articol"], "units": {"Pret_Referinta_RON": "lei"}, "null_default": {"Producator": "Producător local"}, "tests": 1},
                         "materialele (cu producătorul) și serviciile subcontractate"),
        "Dim_Furnizori": (furnizori, {"case": ["Tip_Furnizor"], "trim": ["Furnizor"], "dupes": 2},
                          "producătorii, distribuitorii locali și subcontractorii"),
    }
    return {"tema": "erbasu", "share": share, "tables": tables, "dim_date_desc": "calendarul 2023–2025",
            "relatii": [("Fact_Situatii_Lucrari", "ID_Proiect", "Dim_Proiecte"), ("Fact_Situatii_Lucrari", "ID_Categorie", "Dim_Categorii_Lucrari"),
                        ("Fact_Situatii_Lucrari", "Luna", "Dim_Date"), ("Fact_Achizitii", "ID_Proiect", "Dim_Proiecte"),
                        ("Fact_Achizitii", "ID_Articol", "Dim_Articole"), ("Fact_Achizitii", "ID_Furnizor", "Dim_Furnizori"),
                        ("Fact_Achizitii", "Data_Achizitie", "Dim_Date"), ("Dim_Proiecte", "ID_Beneficiar", "Dim_Beneficiari")]}


# ===========================================================================
# 3. ENDAVA ROMÂNIA
# ===========================================================================
def build_endava() -> dict:
    ref = load_ref("endava")
    rng = random.Random(1203)
    cal = ref["calibrare"]

    jud = {"Cluj-Napoca": "Cluj", "București": "București", "Iași": "Iași", "Timișoara": "Timiș", "Brașov": "Brașov",
           "Sibiu": "Sibiu", "Craiova": "Dolj", "Pitești": "Argeș", "Târgu Mureș": "Mureș", "Suceava": "Suceava"}
    centre = [{"ID_Centru": i, "Oras": o, "Judet": jud[o], "Sediu_Social": "Da" if o == "Cluj-Napoca" else "Nu"}
              for i, o in enumerate(ref["centre_romania"], 1)]

    tari = {"North America": ["United States", "Canada"], "United Kingdom": ["United Kingdom"],
            "Europe": ["Germany", "Netherlands", "Switzerland", "Denmark"], "Rest of World": ["Australia", "Singapore"]}
    clienti = [{"ID_Client": 1, "Client": "Endava (intern)", "Industrie": "Intern", "Geografie": "Romania", "Tara": "Romania"}]
    for c in ref["clienti_publici"]:
        clienti.append({"ID_Client": len(clienti) + 1, "Client": c["nume"], "Industrie": c["industrie"], "Geografie": "United Kingdom",
                        "Tara": c["tara"]})
    counter = {}
    ind_w = cal["mix_industrii_FY2025"]
    for _ in range(26):
        ind, geo = pick(rng, ind_w), pick(rng, cal["mix_geografii_FY2025"])
        tara = rng.choice(tari[geo])
        k = (ind, tara)
        counter[k] = counter.get(k, 0) + 1
        short = {"Banking & Capital Markets": "BCM"}.get(ind, ind)
        code = {"United States": "US", "Canada": "CA", "United Kingdom": "UK", "Germany": "DE", "Netherlands": "NL",
                "Switzerland": "CH", "Denmark": "DK", "Australia": "AU", "Singapore": "SG"}[tara]
        clienti.append({"ID_Client": len(clienti) + 1, "Client": f"Client {short} {code} #{counter[k]}", "Industrie": ind,
                        "Geografie": geo, "Tara": tara})

    proiecte = [{"ID_Proiect": 1, "Proiect": "Bench & formare internă", "ID_Client": 1, "Tip_Proiect": "Intern", "Metodologie": "Kanban",
                 "Data_Start": D0, "Data_Sfarsit_Planificata": D1, "Buget_RON": None, "Status": "Activ", "Facturabil": "Nu"},
                {"ID_Proiect": 2, "Proiect": "Pre-sales & propuneri", "ID_Client": 1, "Tip_Proiect": "Intern", "Metodologie": "Kanban",
                 "Data_Start": D0, "Data_Sfarsit_Planificata": D1, "Buget_RON": None, "Status": "Activ", "Facturabil": "Nu"}]
    for c in clienti[1:]:
        for _ in range(2 if c["Client"] in ("Worldpay", "Paysafe") else rng.choice([1, 1, 2])):
            tip = rng.choice(ref["tipuri_proiect"])
            s = date(2023, 1, 1) + timedelta(rng.randint(-200, 900))
            dur = rng.randint(6, 30)
            e = s + timedelta(dur * 30)
            st = "Finalizat" if e < date(2025, 10, 1) else "Activ"
            if rng.random() < 0.07:
                st = rng.choice(["Anulat", "În pauză"])
            proiecte.append({"ID_Proiect": len(proiecte) + 1, "Proiect": f"{c['Client'].replace('Client ', '')} – {tip}",
                             "ID_Client": c["ID_Client"], "Tip_Proiect": tip, "Metodologie": pick(rng, {"Scrum": 0.55, "Kanban": 0.2, "SAFe": 0.15, "Waterfall": 0.1}),
                             "Data_Start": max(s, date(2021, 1, 1)), "Data_Sfarsit_Planificata": e, "Buget_RON": 0.0,
                             "Status": st, "Facturabil": "Da"})

    rol_w = {r["rol"]: r["pondere"] for r in ref["roluri"]}
    sen = {s["nivel"]: s for s in ref["senioritati"]}
    ang = []
    hc = cal["angajati_medii"]  # 4372, 3333, 4023 → eșantion de 10%
    n0 = round(hc["2023"] / 10 * 1.06)
    for i in range(n0 + 130):
        nivel = pick(rng, {k: v["pondere"] for k, v in sen.items()})
        lo, hi = sen[nivel]["cost_brut_lunar_ron"]
        cost = rng.uniform(lo, hi)
        if i < n0:
            start = date(rng.randint(2014, 2022), rng.randint(1, 12), 1)
        else:
            start = date(2025, rng.randint(1, 9), 1) if i < n0 + 100 else date(2023, rng.randint(1, 12), 1)
        ang.append({"ID_Angajat": i + 1, "Nume": person(rng), "Rol": pick(rng, rol_w), "Senioritate": nivel,
                    "ID_Centru": pick(rng, {c["ID_Centru"]: ref["pondere_centre"][c["Oras"]] for c in centre}),
                    "Cost_Orar_RON": r2(cost * 1.0225 / 168), "Capacitate_Ore_Luna": 168 if rng.random() > 0.04 else 84,
                    "Data_Angajare": start, "Data_Plecare": None})
    # plecările din 2023–2024 (restructurarea raportată: de la 4.372 la 3.333 angajați medii)
    leavers = rng.sample([a for a in ang if a["Data_Angajare"].year < 2025], round(n0 * 0.36))
    for a in leavers:
        a["Data_Plecare"] = date(2023, 6, 1) + timedelta(rng.randint(0, 540))

    def active(a, m):
        return a["Data_Angajare"] <= m and (a["Data_Plecare"] is None or a["Data_Plecare"] > m)

    proj_client = [p for p in proiecte if p["Facturabil"] == "Da"]
    alocari, pontaj = [], []
    asign = {}
    for m in months(D0, D1):
        live = [p for p in proj_client if p["Data_Start"] <= m <= p["Data_Sfarsit_Planificata"] and p["Status"] != "Anulat"]
        for a in ang:
            if not active(a, m):
                continue
            cur = asign.get(a["ID_Angajat"])
            if a["Rol"] == "Support / Bench" or not live:
                plan = [(1, 1.0)]
            else:
                live_ids = {p["ID_Proiect"] for p in live}
                if not cur or any(pid > 2 and pid not in live_ids for pid, _ in cur) or rng.random() < 0.04:
                    p1 = rng.choice(live)
                    x = rng.random()
                    if x < 0.12:  # supra-alocat
                        p2 = rng.choice(live)
                        cur = [(p1["ID_Proiect"], 1.0), (p2["ID_Proiect"], rng.choice([0.2, 0.3, 0.5]))] if p2 is not p1 else [(p1["ID_Proiect"], 1.2)]
                    elif x < 0.22:  # sub-alocat, restul pe bench
                        f = rng.choice([0.4, 0.5, 0.6, 0.8])
                        cur = [(p1["ID_Proiect"], f), (1, round(1 - f, 2))]
                    elif x < 0.3 and m.year == 2025:  # 2025: mai mulți oameni pe bench (venituri în scădere)
                        cur = [(1, 1.0)]
                    else:
                        cur = [(p1["ID_Proiect"], 1.0)]
                    asign[a["ID_Angajat"]] = cur
                plan = cur
            for pid, f in plan:
                ore_plan = round(a["Capacitate_Ore_Luna"] * f)
                alocari.append({"ID_Alocare": 0, "Luna": m, "ID_Angajat": a["ID_Angajat"], "ID_Proiect": pid,
                                "Alocare_Planificata": f, "Ore_Planificate": ore_plan})
                lucrat = ore_plan * rng.uniform(0.82, 1.06)
                if a["Data_Plecare"] and (a["Data_Plecare"] - m).days < 31:
                    lucrat *= 0.5
                fact = pid > 2 and rng.random() > 0.03
                pontaj.append({"ID_Pontaj": 0, "Luna": m, "ID_Angajat": a["ID_Angajat"], "ID_Proiect": pid,
                               "Ore_Lucrate": r2(lucrat), "Ore_Facturabile": r2(lucrat * rng.uniform(0.9, 1.0)) if fact else 0.0,
                               "Ore_Suplimentare": r2(max(0, lucrat - a["Capacitate_Ore_Luna"])) if f >= 1 else 0.0,
                               "Cost_RON": r2(lucrat * a["Cost_Orar_RON"]), "Tarif_Ora_RON": 0.0, "Venit_RON": 0.0})
    # tariful facturat pe an: calibrat pe cifra de afaceri reală × ponderea eșantionului
    ca = {int(k): v for k, v in cal["cifra_afaceri_ron"].items()}
    share = len([a for a in ang if active(a, date(2024, 6, 1))]) / hc["2024"]
    for y in (2023, 2024, 2025):
        ore_f = sum(p["Ore_Facturabile"] for p in pontaj if p["Luna"].year == y)
        base = ca[y] * share / ore_f
        for p in pontaj:
            if p["Luna"].year == y and p["Ore_Facturabile"]:
                p["Tarif_Ora_RON"] = r2(base * rng.uniform(0.85, 1.15))
                p["Venit_RON"] = r2(p["Ore_Facturabile"] * p["Tarif_Ora_RON"])
    for p in proiecte:
        if p["Facturabil"] == "Da":
            v = sum(x["Venit_RON"] for x in pontaj if x["ID_Proiect"] == p["ID_Proiect"])
            p["Buget_RON"] = r2(max(v, 150000) * rng.uniform(0.9, 1.25)) if v else r2(rng.uniform(4e5, 3e6))
    alocari = order(alocari, "ID_Alocare")
    pontaj = order(pontaj, "ID_Pontaj")

    riscuri = []
    for p in proiecte[2:]:
        for _ in range(rng.randint(1, 5)):
            pr, im = rng.randint(1, 5), rng.randint(1, 5)
            ident = max(p["Data_Start"], D0) + timedelta(rng.randint(0, 300))
            if ident > D1:
                ident = D1 - timedelta(rng.randint(5, 60))
            st = pick(rng, {"Deschis": 0.3, "Monitorizat": 0.3, "Închis": 0.3, "Materializat": 0.1})
            riscuri.append({"ID_Risc": len(riscuri) + 1, "ID_Proiect": p["ID_Proiect"], "Categorie_Risc": rng.choice(ref["categorii_risc"]),
                            "Probabilitate": pr, "Impact": im, "Scor_Risc": pr * im, "Status_Risc": st, "Data_Identificare": ident})

    tables = {
        "Fact_Pontaj": (pontaj, {"comma": ["Ore_Lucrate", "Ore_Facturabile", "Cost_RON"], "date_text": ["Luna"], "id_text": ["ID_Angajat"],
                                 "extra": "Exportat_La", "dupes": 9, "tests": 2, "blanks": 3, "total": ["Venit_RON"]},
                         "orele pontate lunar pe angajat și proiect: facturabile, suplimentare, cost, venit"),
        "Fact_Alocari": (alocari, {"pct": ["Alocare_Planificata"], "units": {"Ore_Planificate": "h"}, "date_text": ["Luna"],
                                   "dupes": 6, "blanks": 2, "title": "Planificare resurse - export Resource Management"},
                          "alocarea planificată pe angajat și proiect în fiecare lună (% din capacitate)"),
        "Dim_Angajati": (ang, {"case": ["Rol", "Senioritate"], "trim": ["Nume"], "comma": ["Cost_Orar_RON"], "date_text": ["Data_Angajare"], "dupes": 3, "tests": 1},
                         "angajații din centrele din România: rol, senioritate, cost orar, capacitate, angajare/plecare"),
        "Dim_Proiecte": (proiecte, {"case": ["Status", "Metodologie"], "trim": ["Proiect"], "dupes": 1, "null_default": {"Facturabil": "Nu"}},
                         "proiectele (plus bench și pre-sales): client, tip, metodologie, buget, termene"),
        "Dim_Clienti": (clienti, {"case": ["Industrie"], "trim": ["Client"],
                                  "synonyms": {"Geografie": {"United Kingdom": "UK", "North America": "NA"}}},
                        "clienții: industrie, geografie, țară"),
        "Dim_Centre": (centre, {"trim": ["Oras"], "synonyms": {"Oras": {"Cluj-Napoca": "Cluj Napoca", "Târgu Mureș": "Targu Mures"}}},
                       "cele 10 orașe cu birouri Endava din România"),
        "Dim_Riscuri": (riscuri, {"case": ["Status_Risc", "Categorie_Risc"], "date_text": ["Data_Identificare"], "dupes": 2},
                        "registrul de riscuri pe proiect: probabilitate, impact, scor, status"),
    }
    return {"tema": "endava", "share": share, "tables": tables, "dim_date_desc": "calendarul 2023–2025 (faptele sunt lunare: prima zi a lunii)",
            "relatii": [("Fact_Pontaj", "ID_Angajat", "Dim_Angajati"), ("Fact_Pontaj", "ID_Proiect", "Dim_Proiecte"), ("Fact_Pontaj", "Luna", "Dim_Date"),
                        ("Fact_Alocari", "ID_Angajat", "Dim_Angajati"), ("Fact_Alocari", "ID_Proiect", "Dim_Proiecte"), ("Fact_Alocari", "Luna", "Dim_Date"),
                        ("Dim_Proiecte", "ID_Client", "Dim_Clienti"), ("Dim_Angajati", "ID_Centru", "Dim_Centre"), ("Dim_Riscuri", "ID_Proiect", "Dim_Proiecte")]}


# ===========================================================================
# 4. UNILEVER ROMÂNIA
# ===========================================================================
def build_unilever() -> dict:
    ref = load_ref("unilever")
    rng = random.Random(1204)
    cats = list(ref["calibrare"]["mix_categorii"])
    categorii = [{"ID_Categorie": i, "Categorie": c, "Divizie": {"Home Care": "Home Care", "Beauty & Personal Care": "Beauty & Wellbeing",
                  "Nutrition (Foods)": "Foods", "Ice Cream": "Ice Cream (TMICC din dec. 2025)"}[c]} for i, c in enumerate(cats, 1)]
    cat_id = {c["Categorie"]: c["ID_Categorie"] for c in categorii}
    brand = {b["brand"]: b for b in ref["branduri"]}
    produse = []
    for i, s in enumerate(ref["sku"], 1):
        b = brand[s["brand"]]
        produse.append({"ID_Produs": i, "Produs": s["produs"], "Brand": s["brand"], "ID_Categorie": cat_id[b["categorie"]],
                        "Subcategorie": b["subcategorie"], "Gramaj": s["gramaj"], "Bucati_Bax": s["buc_bax"],
                        "Fabrica": b.get("fabrica", "Import grup Unilever"), "Pret_Raft_RON": s["pret_raft_ron"]})

    zone = ref["zone_kam"]
    agenti = []
    for z in zone:
        for _ in range(3 if z == "Key Accounts Național" else 2):
            agenti.append({"ID_Agent": len(agenti) + 1, "Nume": person(rng), "Functie": "Key Account Manager" if z == "Key Accounts Național" else "Area Sales Manager",
                           "Zona": z})
    oras = {"Kaufland România": "București", "Lidl România": "Chiajna", "Carrefour România": "București", "Mega Image": "București",
            "Profi": "Timișoara", "Penny": "Ștefăneștii de Jos", "Auchan România": "București", "Selgros": "Brașov", "Metro România": "Voluntari",
            "dm drogerie markt": "București", "eMAG": "Voluntari", "Freshful by eMAG": "Chiajna"}
    pondere = {"Kaufland România": 0.17, "Lidl România": 0.13, "Carrefour România": 0.10, "Mega Image": 0.09, "Profi": 0.08, "Penny": 0.06,
               "Auchan România": 0.05, "Selgros": 0.04, "Metro România": 0.04, "dm drogerie markt": 0.03, "eMAG": 0.02, "Freshful by eMAG": 0.015,
               "Distribuitor HoReCa național": 0.035}
    clienti = []
    for c in ref["clienti"]:
        n = c["client"]
        if n.startswith("Distribuitor TT"):
            z = n.replace("Distribuitor TT ", "")
            ag = rng.choice([a for a in agenti if a["Zona"] == z])
            o = {"București-Ilfov": "București", "Moldova": "Iași", "Transilvania": "Cluj-Napoca", "Muntenia-Oltenia": "Craiova",
                 "Banat-Crișana": "Timișoara", "Dobrogea": "Constanța"}[z]
            w = 0.025
        else:
            ag = rng.choice([a for a in agenti if a["Zona"] == "Key Accounts Național"])
            o, w = oras.get(n, "București"), pondere[n]
        clienti.append({"ID_Client": len(clienti) + 1, "Client": n, "Canal": c["canal"], "Format": c["format"], "Oras_Sediu": o,
                        "ID_Agent": ag["ID_Agent"], "_w": w})
    promotii = [{"ID_Promotie": i, "Promotie": p, "Tip_Promotie": "Fără" if i == 1 else "Sezonieră" if p in ("Black Friday", "Paște", "Crăciun", "Back to School", "Sezon înghețată") else "Comercială",
                 "Discount_Tipic": [0, 0.15, 0.25, 0.10, 0.08, 0.20, 0.12, 0.15, 0.10, 0.12][i - 1]} for i, p in enumerate(ref["promotii"], 1)]

    cat_season = {"Home Care": [1] * 12, "Beauty & Personal Care": [0.9, 0.9, 1, 1, 1.1, 1.2, 1.2, 1.1, 1, 0.95, 1, 1.15],
                  "Nutrition (Foods)": [1.15, 1.05, 1.1, 1.1, 0.9, 0.8, 0.75, 0.8, 0.95, 1.1, 1.2, 1.3],
                  "Ice Cream": [0.15, 0.2, 0.4, 0.8, 1.6, 2.6, 3.0, 2.7, 1.2, 0.4, 0.15, 0.15]}
    infl = {2023: 1.0, 2024: 1.06, 2025: 1.10}
    mix = ref["calibrare"]["mix_categorii"]
    cmap = {c: [p for p in produse if p["ID_Categorie"] == cat_id[c]] for c in cats}

    def line(year):
        c = rng.choices(clienti, weights=[x["_w"] for x in clienti])[0]
        cat = pick(rng, mix)
        if c["Client"] == "dm drogerie markt" and cat not in ("Home Care", "Beauty & Personal Care"):
            return None
        if c["Canal"] == "HoReCa" and cat not in ("Nutrition (Foods)", "Ice Cream"):
            return None
        d = rand_day(rng, year, cat_season[cat])
        if cat == "Ice Cream" and d >= date(2025, 12, 1):
            return None  # înghețata a trecut la The Magnum Ice Cream Company
        if d.weekday() == 6:
            return None
        p = rng.choice(cmap[cat])
        bax = max(1, round(rng.lognormvariate(math.log(40 if c["_w"] > 0.05 else 18), 0.7)))
        pret = r2(p["Pret_Raft_RON"] / 1.19 / 1.25 * infl[year])
        promo = 1
        if rng.random() < 0.28:
            opts = [2, 3, 4, 5]
            m = d.month
            if m == 11: opts.append(6)
            if m in (3, 4): opts.append(7)
            if m == 12: opts.append(8)
            if m == 9: opts.append(9)
            if cat == "Ice Cream" and m in (5, 6, 7, 8): opts += [10, 10]
            promo = rng.choice(opts)
        disc = promotii[promo - 1]["Discount_Tipic"]
        buc = bax * p["Bucati_Bax"]
        return {"ID_Livrare": 0, "Data_Livrare": d, "ID_Client": c["ID_Client"], "ID_Produs": p["ID_Produs"], "ID_Agent": c["ID_Agent"],
                "ID_Promotie": promo, "Baxuri": bax, "Bucati": buc, "Pret_Unitar_RON": pret, "Discount": disc,
                "Valoare_Neta_RON": r2(buc * pret * (1 - disc))}

    trend = {int(k): v for k, v in ref["calibrare"]["cifra_afaceri_ron"].items()}
    vanzari, share = calibrate(rng, trend, line, "Valoare_Neta_RON", 14000)
    vanzari = order(vanzari, "ID_Livrare")

    pcat = {p["ID_Produs"]: p["ID_Categorie"] for p in produse}
    real = {}
    for v in vanzari:
        k = (v["ID_Agent"], pcat[v["ID_Produs"]], date(v["Data_Livrare"].year, v["Data_Livrare"].month, 1))
        real[k] = real.get(k, 0) + v["Valoare_Neta_RON"]
    target = []
    for (ag, cid, m), val in sorted(real.items(), key=lambda kv: (kv[0][2], kv[0][0], kv[0][1])):
        target.append({"ID_Target": len(target) + 1, "Luna": m, "ID_Agent": ag, "ID_Categorie": cid,
                       "Target_RON": r2(round(val * rng.uniform(0.9, 1.14), -2))})
    for c in clienti:
        del c["_w"]

    tables = {
        "Fact_Vanzari": (vanzari, {"date_text": ["Data_Livrare"], "units": {"Baxuri": "bax"}, "comma": ["Pret_Unitar_RON"], "pct": ["Discount"],
                                   "id_text": ["ID_Produs"], "extra": "Exportat_La", "dupes": 9, "tests": 2, "blanks": 3, "total": ["Valoare_Neta_RON"]},
                         "livrările (sell-in) pe client și produs: baxuri, bucăți, preț, promoție, valoare netă"),
        "Fact_Target": (target, {"date_text": ["Luna"], "comma": ["Target_RON"], "dupes": 4, "blanks": 1, "title": "Targete lunare vânzări 2023-2025"},
                        "targetul lunar pe agent și categorie"),
        "Dim_Produse": (produse, {"trim": ["Produs"], "case": ["Brand"], "units": {"Bucati_Bax": "buc"}, "comma": ["Pret_Raft_RON"], "dupes": 1, "tests": 1},
                        "produsele (SKU): brand, gramaj, bucăți pe bax, fabrică, preț de raft"),
        "Dim_Categorii": (categorii, {"trim": ["Categorie"]}, "cele 4 categorii de business Unilever"),
        "Dim_Clienti": (clienti, {"case": ["Canal"], "trim": ["Client"], "synonyms": {"Canal": {"Modern Trade": "MT", "Traditional Trade": "TT"}}, "dupes": 1},
                        "lanțurile de retail, cash & carry, online și distribuitorii"),
        "Dim_Agenti": (agenti, {"case": ["Functie"], "trim": ["Nume"]}, "KAM-ii și managerii de zonă"),
        "Dim_Promotii": (promotii, {"pct": ["Discount_Tipic"], "case": ["Tip_Promotie"]}, "mecanismele promoționale și discountul tipic"),
    }
    return {"tema": "unilever", "share": share, "tables": tables, "dim_date_desc": "calendarul 2023–2025",
            "relatii": [("Fact_Vanzari", "ID_Client", "Dim_Clienti"), ("Fact_Vanzari", "ID_Produs", "Dim_Produse"), ("Fact_Vanzari", "ID_Agent", "Dim_Agenti"),
                        ("Fact_Vanzari", "ID_Promotie", "Dim_Promotii"), ("Fact_Vanzari", "Data_Livrare", "Dim_Date"),
                        ("Fact_Target", "ID_Agent", "Dim_Agenti"), ("Fact_Target", "ID_Categorie", "Dim_Categorii"), ("Fact_Target", "Luna", "Dim_Date"),
                        ("Dim_Produse", "ID_Categorie", "Dim_Categorii")]}


# ===========================================================================
# 5. UIPATH
# ===========================================================================
UIP_ORE_AN = {"Royal Mail": 130000, "Mediq": 55000, "Bank ABB": 34000, "Sonae Arauco": 8000, "City of Copenhagen": 8500,
              "Australian Unity": 9000, "The New York Foundling": 100000, "Banca Transilvania": 40000,
              "Provident Financial Romania": 10000, "Major Members-only Retailer (via Accelirate)": 1920,
              "Evros Technology Group": 1100, "Canon USA": 1800}
UIP_PROC = {"Logistică și poștă": [12, 1, 9, 10, 5, 11], "Distribuție medicală": [1, 9, 12, 10, 2], "Bancar": [6, 7, 8, 2, 1, 11, 4],
            "Producție (panouri lemn)": [3, 1, 9, 5], "Administrație publică": [1, 4, 5, 10], "Asigurări și sănătate": [7, 10, 1, 4, 14],
            "ONG / servicii sociale": [4, 5, 1, 10], "Servicii financiare": [4, 5, 6, 2, 3], "Retail": [1], "Servicii IT": [1], "Echipamente electronice": [1, 9]}


def build_uipath() -> dict:
    ref = load_ref("uipath")
    rng = random.Random(1205)
    dep = [{"ID_Departament": i, "Departament": d} for i, d in enumerate(ref["departamente"], 1)]
    dep_id = {d["Departament"]: d["ID_Departament"] for d in dep}
    procese = []
    for i, p in enumerate(ref["procese_tip"], 1):
        procese.append({"ID_Proces": i, "Proces": p["proces"], "ID_Departament": dep_id[p["departament"]],
                        "Durata_Manual_Min": r2(sum(p["durata_manual_min"]) / 2), "Durata_Automat_Min": r2(sum(p["durata_automat_min"]) / 2),
                        "Complexitate": "Ridicată" if p["durata_manual_min"][1] > 100 else "Medie" if p["durata_manual_min"][1] > 15 else "Scăzută",
                        "_p": p})
    produse = [{"ID_Produs": i, "Produs": p["produs"], "Familie": p["familie"], "Tip": p["tip"],
                "Disponibil_Din": date(2025, 1, 1) if p["familie"] == "Agentic" or p["produs"] == "Autopilot" else date(2018, 1, 1)}
               for i, p in enumerate(ref["produse_uipath"], 1)]
    pid = {p["Produs"]: p["ID_Produs"] for p in produse}
    clienti = [{"ID_Client": i, "Client": c["client"], "Tara": c["tara"], "Industrie": c["industrie"],
                "Client_Din": date(2017, 6, 1) if c["client"] == "Banca Transilvania" else date(rng.randint(2018, 2022), rng.randint(1, 12), 1)}
               for i, c in enumerate(ref["clienti_studii_caz"], 1)]
    cost_h = ref["cost_orar_operator_usd"]

    initiative = []
    for c in clienti:
        lst = UIP_PROC[c["Industrie"]]
        for k, pr_id in enumerate(lst):
            p = procese[pr_id - 1]
            # go-live: unele procese erau automatizate înainte de 2023, altele intră în perioada setului
            if k < len(lst) // 2 and c["Client_Din"].year < 2021:
                gl = date(rng.randint(2019, 2022), rng.randint(1, 12), 1)
            else:
                gl = date(2023, 3, 1) + timedelta(rng.randint(0, 900))
            st = date(gl.year, gl.month, 1) - timedelta(rng.randint(60, 200))
            status = "Live" if gl <= D1 else "În dezvoltare"
            if gl > date(2025, 9, 1) and rng.random() < 0.4:
                status = "Pilot"
            if rng.random() < 0.05:
                status = "Oprit"
            prod = ("Document Understanding" if pr_id in (1, 9) else "Communications Mining" if pr_id == 10 else "Test Suite" if pr_id == 14
                    else "Agent Builder" if gl >= date(2025, 3, 1) and rng.random() < 0.35 else "Unattended Robot")
            initiative.append({"ID_Initiativa": len(initiative) + 1, "ID_Client": c["ID_Client"], "ID_Proces": pr_id, "ID_Produs": pid[prod],
                               "Data_Start": st, "Data_Go_Live": gl, "Status": status, "Nr_Roboti": rng.randint(1, 6),
                               "Investitie_USD": r2(rng.uniform(25000, 180000) * (1.6 if p["Complexitate"] == "Ridicată" else 1)),
                               "Ore_Economisite_Estimate_An": 0.0})
    executii = []
    pmap = {p["ID_Proces"]: p for p in procese}
    for ini in initiative:
        c = clienti[ini["ID_Client"] - 1]
        p = pmap[ini["ID_Proces"]]["_p"]
        n_proc = len(UIP_PROC[c["Industrie"]])
        man = sum(p["durata_manual_min"]) / 2
        aut = sum(p["durata_automat_min"]) / 2
        saved_h = (man - aut) / 60
        vol_month = max(20, round(UIP_ORE_AN[c["Client"]] / n_proc / 12 / saved_h))
        ini["Ore_Economisite_Estimate_An"] = r2(vol_month * 12 * saved_h * rng.uniform(1.0, 1.3))
        ch = cost_h[{"United Kingdom": "United Kingdom", "United States": "United States"}.get(c["Tara"], c["Tara"])] if c["Tara"] in cost_h else 30
        for m in months(D0, D1):
            vol = round(vol_month * rng.uniform(0.85, 1.15) * (1.08 if m.month in (1, 3, 12) else 1))
            live = ini["Status"] in ("Live", "Oprit") and ini["Data_Go_Live"] <= m and not (ini["Status"] == "Oprit" and m >= date(2025, 6, 1))
            if live:
                months_live = (m.year - ini["Data_Go_Live"].year) * 12 + m.month - ini["Data_Go_Live"].month
                ramp = min(1, 0.55 + 0.15 * months_live)
                auto_share = ramp * rng.uniform(0.86, 0.97)
                parts = [("Automat", round(vol * auto_share)), ("Manual", vol - round(vol * auto_share))]
            else:
                parts = [("Manual", vol)]
            for mod, n in parts:
                if n <= 0:
                    continue
                dur = man if mod == "Manual" else aut
                rate = p["rata_erori_manual"] if mod == "Manual" else p["rata_erori_automat"]
                tot = n * dur * rng.uniform(0.9, 1.12)
                err = sum(1 for _ in range(min(n, 400)) if rng.random() < rate) * (n / min(n, 400))
                err = round(err)
                cost = tot / 60 * ch if mod == "Manual" else n * rng.uniform(0.04, 0.25) + err * dur / 60 * ch
                executii.append({"ID_Executie": 0, "Luna": m, "ID_Client": c["ID_Client"], "ID_Proces": ini["ID_Proces"],
                                 "ID_Initiativa": ini["ID_Initiativa"], "Mod_Executie": mod, "Nr_Executii": n,
                                 "Durata_Totala_Min": r2(tot), "Nr_Erori": err, "Nr_Rework": round(err * rng.uniform(0.4, 0.8)),
                                 "Cost_USD": r2(cost)})
    executii = order(executii, "ID_Executie")
    for p in procese:
        del p["_p"]
    # Fact_Initiative are ID_Initiativa, dar relaționarea se face prin client/proces/produs/dată: nu între fapte
    for e in executii:
        del e["ID_Initiativa"]

    tables = {
        "Fact_Executii": (executii, {"date_text": ["Luna"], "comma": ["Durata_Totala_Min", "Cost_USD"], "case": ["Mod_Executie"],
                                     "units": {"Nr_Executii": "exec."}, "id_text": ["ID_Proces"], "extra": "Exportat_La",
                                     "dupes": 8, "tests": 2, "blanks": 3, "total": ["Cost_USD"],
                                     "synonyms": {"Mod_Executie": {"Automat": "Robot", "Manual": "Uman"}}},
                           "execuțiile lunare pe client și proces, Manual vs. Automat: număr, durată, erori, rework, cost"),
        "Fact_Initiative": (initiative, {"date_text": ["Data_Go_Live"], "comma": ["Investitie_USD", "Ore_Economisite_Estimate_An"], "case": ["Status"],
                                         "dupes": 2, "blanks": 1, "title": "Portofoliu automatizări - export CoE"},
                            "proiectele de automatizare: produs folosit, start, go-live, status, investiție, ore economisite estimate"),
        "Dim_Clienti": (clienti, {"trim": ["Client"], "case": ["Industrie"], "date_text": ["Client_Din"], "tests": 1},
                        "clienții UiPath: țară, industrie, de când sunt clienți"),
        "Dim_Procese": (procese, {"case": ["Complexitate"], "comma": ["Durata_Manual_Min", "Durata_Automat_Min"], "trim": ["Proces"], "dupes": 1},
                        "procesele automatizate: departament, durata manuală vs. automată, complexitate"),
        "Dim_Produse": (produse, {"case": ["Familie"], "date_text": ["Disponibil_Din"]},
                        "produsele platformei UiPath (RPA, IDP, Process Mining, Agentic)"),
        "Dim_Departamente": (dep, {"trim": ["Departament"]}, "departamentele în care rulează procesele"),
    }
    return {"tema": "uipath", "share": None, "tables": tables, "dim_date_desc": "calendarul 2023–2025 (Fact_Executii e lunar)",
            "relatii": [("Fact_Executii", "ID_Client", "Dim_Clienti"), ("Fact_Executii", "ID_Proces", "Dim_Procese"), ("Fact_Executii", "Luna", "Dim_Date"),
                        ("Fact_Initiative", "ID_Client", "Dim_Clienti"), ("Fact_Initiative", "ID_Proces", "Dim_Procese"), ("Fact_Initiative", "ID_Produs", "Dim_Produse"),
                        ("Dim_Procese", "ID_Departament", "Dim_Departamente")]}


# ===========================================================================
# 6. H&M ROMÂNIA
# ===========================================================================
DENUMIRI = {"T-shirt": "Tricou", "Top": "Top", "Blouse": "Bluză", "Shirt": "Cămașă", "Sweater": "Pulover", "Cardigan": "Cardigan",
            "Hoodie": "Hanorac", "Jacket": "Jachetă", "Coat": "Palton", "Blazer": "Sacou", "Trousers": "Pantaloni", "Jeans": "Blugi",
            "Shorts": "Pantaloni scurți", "Skirt": "Fustă", "Leggings/Tights": "Colanți", "Dress": "Rochie", "Jumpsuit/Playsuit": "Salopetă",
            "Bra": "Sutien", "Underwear bottom": "Chiloți", "Socks": "Șosete", "Pyjama set": "Pijama", "Bikini top": "Sutien de baie",
            "Swimwear bottom": "Slip de baie", "Swimsuit": "Costum de baie", "Sneakers": "Sneakers", "Boots": "Ghete", "Sandals": "Sandale",
            "Bag": "Geantă", "Hat/beanie": "Căciulă", "Scarf": "Fular", "Earring": "Cercei", "Sunglasses": "Ochelari de soare", "Belt": "Curea"}
CULORI = {"Black": "Negru", "Dark Blue": "Bleumarin", "White": "Alb", "Light Pink": "Roz deschis", "Grey": "Gri", "Light Beige": "Bej deschis",
          "Blue": "Albastru", "Beige": "Bej", "Off White": "Alb unt", "Dark Green": "Verde închis", "Red": "Roșu", "Light Blue": "Bleu",
          "Dark Grey": "Gri închis", "Brown": "Maro", "Other": "Multicolor"}
SEZON_TIP = {"Coat": "FW", "Sweater": "FW", "Cardigan": "FW", "Boots": "FW", "Hat/beanie": "FW", "Scarf": "FW", "Jacket": "FW",
             "Shorts": "SS", "Sandals": "SS", "Bikini top": "SS", "Swimwear bottom": "SS", "Swimsuit": "SS", "Sunglasses": "SS"}


def build_hm() -> dict:
    ref = load_ref("hm")
    rng = random.Random(1206)
    tax = ref["taxonomie_articole"]
    preturi = ref["preturi_ron"]
    magazine = [{"ID_Magazin": i, "Magazin": m["magazin"], "Oras": m["oras"], "Judet": m["judet"], "Adresa": m["adresa"],
                 "Canal": "Online" if m["oras"] == "Online" else "Magazin fizic"} for i, m in enumerate(ref["magazine"], 1)]
    online_id = magazine[-1]["ID_Magazin"]
    w_mag = {m["ID_Magazin"]: (1.6 if m["Magazin"] in ("AFI Cotroceni", "Băneasa Shopping City", "Mega Mall", "Iulius Mall Cluj", "Palas Mall", "Iulius Town Timișoara") else 1.0)
             for m in magazine[:-1]}

    pg_of = {t: g for g, ts in tax["product_group"].items() for t in ts}
    colectii = [f"{s}{y}" for y in (22, 23, 24, 25) for s in ("SS", "FW")]
    produse = []
    for i in range(1, 521):
        dep = pick(rng, tax["index_group"])
        tip = rng.choice([t for t in pg_of if not (dep == "Menswear" and t in ("Dress", "Skirt", "Bra", "Bikini top", "Blouse", "Earring"))])
        sez = SEZON_TIP.get(tip) or rng.choice(["SS", "FW"])
        an = rng.choice([23, 24, 25, 25] if sez else [23, 24, 25])
        col = pick(rng, tax["colour_group"])
        lo, hi = preturi[tip]
        pret = round(rng.uniform(lo, hi) / 10) * 10 - 0.01
        if dep == "Baby/Children":
            pret = round(pret * 0.65 / 5) * 5 - 0.01
        produse.append({"ID_Produs": i, "Cod_Articol": f"{rng.randint(500000, 999999):06d}{rng.randint(1, 9):03d}",
                        "Denumire": f"{DENUMIRI[tip]} {rng.choice(['basic', 'regular fit', 'oversized', 'slim fit', 'din bumbac', 'cu imprimeu', 'relaxed fit', 'din in', 'tricotat'])}",
                        "Tip_Produs": DENUMIRI[tip], "Grup_Produs": pg_of[tip], "Departament": dep, "Culoare": CULORI[col],
                        "Model_Grafic": rng.choice(tax["graphical_appearance"]), "Colectie": f"{sez}{an}", "Sezon": sez,
                        "Pret_Lista_RON": max(14.99, pret)})

    v_w = ref["clienti"]["varsta"]
    clienti = [{"ID_Client": 0, "Cod_Client": "NEMEMBRU", "Grupa_Varsta": None, "Varsta": None, "Judet": None,
                "Status_Membru": "Fără card", "Newsletter": None, "Data_Inscriere": None}]
    judete = [m["Judet"] for m in magazine[:-1]]
    for i in range(1, 3201):
        g = pick(rng, v_w)
        lo, hi = (16, 24) if g == "16-24" else (55, 75) if g == "55+" else (int(g[:2]), int(g[3:]))
        clienti.append({"ID_Client": i, "Cod_Client": "%012x" % rng.getrandbits(48), "Grupa_Varsta": g, "Varsta": rng.randint(lo, hi),
                        "Judet": rng.choice(judete), "Status_Membru": pick(rng, ref["clienti"]["club_member_status"]),
                        "Newsletter": pick(rng, ref["clienti"]["fashion_news_frequency"]),
                        "Data_Inscriere": date(2019, 1, 1) + timedelta(rng.randint(0, 2400))})
    campanii = [{"ID_Campanie": i, "Campanie": c["campanie"], "Luni": ", ".join(LUNI[m - 1] for m in c["luni"]) or "-",
                 "Reducere_Min": c["reducere"][0], "Reducere_Max": c["reducere"][1]} for i, c in enumerate(ref["campanii"], 1)]
    fara = campanii[-1]["ID_Campanie"]
    month_w = [0.85, 0.7, 0.85, 0.9, 0.95, 0.9, 0.95, 0.9, 1.0, 1.0, 1.25, 1.55]
    sizes_ad = ["XS", "S", "M", "L", "XL"]
    sizes_kid = ["92", "98", "104", "110", "116", "122", "128", "134", "140"]
    cnt = [0]

    def bon(year):
        d = rand_day(rng, year, month_w)
        onl = rng.random() < 0.18
        mag = online_id if onl else pick(rng, w_mag)
        member = rng.random() < (0.95 if onl else 0.55)
        cl = rng.randint(1, 3200) if member else 0
        cnt[0] += 1
        lines = []
        for _ in range(rng.choices([1, 2, 3, 4], [0.45, 0.3, 0.15, 0.1])[0]):
            p = rng.choice(produse)
            sez = p["Sezon"]
            ok_season = (sez == "SS") == (d.month in ref["sezoane"]["SS"])
            if not ok_season and rng.random() < 0.7:
                p = rng.choice(produse)
            camp = fara
            red = 0.0
            for c in ref["campanii"][:-1]:
                if d.month in c["luni"] and rng.random() < 0.45:
                    camp = campanii[ref["campanii"].index(c)]["ID_Campanie"]
                    red = round(rng.uniform(*c["reducere"]), 1)
                    break
            q = rng.choices([1, 2, 3], [0.85, 0.12, 0.03])[0]
            size = rng.choice(sizes_kid if p["Departament"] == "Baby/Children" else sizes_ad if p["Grup_Produs"] not in ("Shoes", "Accessories") else ["Unică"] if p["Grup_Produs"] == "Accessories" else ["37", "38", "39", "40", "41", "42", "43"])
            lines.append({"ID_Linie": 0, "ID_Bon": f"B{year % 100}{cnt[0]:06d}", "Data_Vanzare": d, "ID_Magazin": mag, "ID_Client": cl,
                          "ID_Produs": p["ID_Produs"], "ID_Campanie": camp, "Marime": size, "Cantitate": q,
                          "Pret_Lista_RON": p["Pret_Lista_RON"], "Reducere": red, "Valoare_RON": r2(q * p["Pret_Lista_RON"] * (1 - red))})
        return lines

    def as_line(year):
        ls = bon(year)
        return {"_lines": ls, "v": sum(x["Valoare_RON"] for x in ls)}

    ca = ref["calibrare"]["cifra_afaceri_ron"]
    trend = {2023: ca["2023"], 2024: ca["2023"] * 1.08, 2025: ca["2023"] * 1.08 * 1.05}
    groups, share = calibrate(rng, trend, as_line, "v", 7200)
    vanzari = [ln for g in groups for ln in g["_lines"]]
    vanzari = order(vanzari, "ID_Linie")

    motive = ref["retururi"]["motive"]
    retururi = []
    for v in vanzari:
        onl = v["ID_Magazin"] == online_id
        if rng.random() > (ref["retururi"]["rata_online"] if onl else ref["retururi"]["rata_magazin"]):
            continue
        z = rng.randint(1, 30)
        dr = v["Data_Vanzare"] + timedelta(z)
        if dr > D1:
            continue
        mot = rng.choices(motive, [0.38, 0.17, 0.12, 0.07, 0.18, 0.05 if onl else 0, 0.03 if onl else 0])[0]
        q = rng.randint(1, v["Cantitate"])
        retururi.append({"ID_Retur": 0, "Data_Retur": dr, "ID_Linie_Vanzare": v["ID_Linie"], "ID_Magazin": v["ID_Magazin"] if not onl or rng.random() < 0.6 else pick(rng, w_mag),
                         "ID_Produs": v["ID_Produs"], "ID_Client": v["ID_Client"], "Canal_Retur": "Curier" if onl and rng.random() < 0.6 else "Magazin",
                         "Motiv_Retur": mot, "Cantitate": q, "Valoare_Returnata_RON": r2(v["Valoare_RON"] / v["Cantitate"] * q), "Zile_Pana_La_Retur": z})
    retururi = order(retururi, "ID_Retur")
    for r in retururi:
        if r["Canal_Retur"] == "Curier":
            r["ID_Magazin"] = online_id

    tables = {
        "Fact_Vanzari": (vanzari, {"date_text": ["Data_Vanzare"], "comma": ["Pret_Lista_RON"], "pct": ["Reducere"], "case": ["Marime"],
                                   "units": {"Cantitate": "buc"}, "id_text": ["ID_Produs"], "extra": "Exportat_La",
                                   "dupes": 9, "tests": 2, "blanks": 3, "total": ["Valoare_RON"]},
                         "liniile de bon din magazine și online: produs, mărime, cantitate, preț, reducere, campanie"),
        "Fact_Retururi": (retururi, {"date_text": ["Data_Retur"], "comma": ["Valoare_Returnata_RON"], "units": {"Zile_Pana_La_Retur": "zile"},
                                     "case": ["Motiv_Retur", "Canal_Retur"], "dupes": 5, "blanks": 2, "title": "Raport retururi H&M Romania"},
                          "retururile: motiv, canal, cantitate, valoare, zile până la retur"),
        "Dim_Produse": (produse, {"case": ["Departament", "Culoare"], "trim": ["Denumire"], "comma": ["Pret_Lista_RON"], "dupes": 3, "tests": 1,
                                  "synonyms": {"Departament": {"Baby/Children": "Kids", "Ladieswear": "Ladies"}}},
                        "articolele: tip, grup de produs, departament, culoare, colecție, preț"),
        "Dim_Magazine": (magazine, {"trim": ["Magazin"], "case": ["Canal"], "synonyms": {"Oras": {"București": "Bucuresti", "Cluj-Napoca": "Cluj Napoca"}}, "dupes": 1},
                         "magazinele H&M (cu adresă) și magazinul online"),
        "Dim_Clienti": (clienti, {"case": ["Status_Membru", "Newsletter"], "null_default": {"Newsletter": "NONE"}, "date_text": ["Data_Inscriere"], "dupes": 4},
                        "membrii H&M și clientul generic fără card (ID 0)"),
        "Dim_Campanii": (campanii, {"pct": ["Reducere_Min", "Reducere_Max"], "trim": ["Campanie"]}, "campaniile de reduceri și lunile în care rulează"),
    }
    return {"tema": "hm", "share": share, "tables": tables, "dim_date_desc": "calendarul 2023–2025",
            "relatii": [("Fact_Vanzari", "ID_Magazin", "Dim_Magazine"), ("Fact_Vanzari", "ID_Client", "Dim_Clienti"), ("Fact_Vanzari", "ID_Produs", "Dim_Produse"),
                        ("Fact_Vanzari", "ID_Campanie", "Dim_Campanii"), ("Fact_Vanzari", "Data_Vanzare", "Dim_Date"),
                        ("Fact_Retururi", "ID_Magazin", "Dim_Magazine"), ("Fact_Retururi", "ID_Client", "Dim_Clienti"),
                        ("Fact_Retururi", "ID_Produs", "Dim_Produse"), ("Fact_Retururi", "Data_Retur", "Dim_Date")]}


# ---------------------------------------------------------------------------
# Rulare
# ---------------------------------------------------------------------------
BUILDERS = [build_telekom, build_erbasu, build_endava, build_unilever, build_uipath, build_hm]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    import pagini_s12 as pg  # noqa: E402

    dim_date = gen_dim_date()
    results = []
    for k, b in enumerate(BUILDERS):
        res = b()
        tema = res["tema"]
        ref = load_ref(tema)
        folder = OUT / tema / "date"
        if folder.exists():
            for f in folder.glob("*.xlsx"):
                f.unlink()
        summary = []
        for j, (name, (rows, spec, desc)) in enumerate(res["tables"].items()):
            drows, info = dirty(rows, spec, 9000 + k * 100 + j)
            check_table(rows, drows, spec)
            write_xlsx(folder / f"{name}.xlsx", name, drows)
            summary.append({"tabel": name, "desc": desc, "randuri": len(rows), "coloane": len(rows[0]),
                            "randuri_murdar": len(drows) - (2 if spec.get("title") else 1), "coloane_murdar": info["cols_dirty"],
                            "murdarie": info["info"], "fapt": name.startswith("Fact")})
        write_xlsx(folder / "Dim_Date.xlsx", "Dim_Date", [list(dim_date[0]), *[list(r.values()) for r in dim_date]])
        summary.append({"tabel": "Dim_Date", "desc": res["dim_date_desc"], "randuri": len(dim_date), "coloane": len(dim_date[0]),
                        "randuri_murdar": len(dim_date), "coloane_murdar": len(dim_date[0]), "murdarie": ["curată, o primiți gata făcută"], "fapt": False})
        res["summary"] = summary
        res["ref"] = ref
        (OUT / tema / "SURSE.txt").unlink(missing_ok=True)
        (OUT / tema / f"date_reale_{tema}.json").unlink(missing_ok=True)
        (OUT / tema / "CITESTE_MAI_INTAI.txt").write_text(pg.readme_txt(tema, summary), encoding="utf-8")
        (OUT / tema / "index.html").write_text(pg.page(tema, ref, summary), encoding="utf-8")
        results.append(res)
        tot = sum(s["randuri"] for s in summary)
        print(f"{tema:9s} {len(summary)} tabele, {tot:6d} rânduri curate, eșantion {res['share'] or 0:.4f} · check OK")
    (OUT / "index.html").write_text(pg.hub(results), encoding="utf-8")
    KEY_OUT.write_text(pg.solutii_md(results), encoding="utf-8")
    print("Gata:", OUT)


if __name__ == "__main__":
    main()
