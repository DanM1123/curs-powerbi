# -*- coding: utf-8 -*-
"""Generează datele pentru partea „Imagini în Power BI” din sesiunea 12.

Dim_Filme: 26 de filme reale. An, regizor, durată, buget și încasări mondiale
sunt luate din infobox-urile Wikipedia, octombrie 2026 (rotunjite, milioane USD;
la intervale, capătul de jos). Nota IMDb e aproximativă. Afișele sunt
link-urile directe către fișierele de pe upload.wikimedia.org.
Fact_Proiectii: programul fictiv al festivalului „Retro & Blockbuster” de la
Cinema Lumière (cinematograf inventat), cu bilete vândute generate.

Scrie:
  sesiuni/s12/date/Dim_Filme.xlsx, Dim_Filme.csv
  sesiuni/s12/date/Fact_Proiectii.xlsx
  js/s12-date.js (aceleași date, pentru demo-ul din slide-uri)
"""
from __future__ import annotations

import csv
import json
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "sesiuni" / "s12" / "date"
WIKI = "https://upload.wikimedia.org/wikipedia/en/"

# ID, titlu, titlu RO, an, regizor, gen, durată, buget, încasări, IMDb, Oscaruri, afiș, descriere
FILME = [
    ("F01", "The Shawshank Redemption", "Închisoarea îngerilor", 1994, "Frank Darabont", "Dramă", 142, 25, 73, 9.3, 0, "8/81/ShawshankRedemptionMoviePoster.jpg",
     "Un bancher condamnat pe nedrept își păstrează speranța în cei aproape 20 de ani petrecuți în închisoarea Shawshank."),
    ("F02", "The Godfather", "Nașul", 1972, "Francis Ford Coppola", "Crimă", 175, 6.5, 250, 9.2, 3, "1/1c/Godfather_ver1.jpg",
     "Familia Corleone și fiul care, fără să vrea, ajunge la conducerea imperiului mafiot."),
    ("F03", "The Dark Knight", "Cavalerul negru", 2008, "Christopher Nolan", "Acțiune", 152, 185, 1009, 9.0, 2, "1/1c/The_Dark_Knight_%282008_film%29.jpg",
     "Batman se confruntă cu Joker, un criminal care vrea să arate că oricine poate fi corupt."),
    ("F04", "Pulp Fiction", "Pulp Fiction", 1994, "Quentin Tarantino", "Crimă", 154, 8.5, 214, 8.9, 1, "3/3b/Pulp_Fiction_%281994%29_poster.jpg",
     "Povești de gangsteri din Los Angeles, spuse în altă ordine decât s-au petrecut."),
    ("F05", "Forrest Gump", "Forrest Gump", 1994, "Robert Zemeckis", "Dramă", 142, 55, 678, 8.8, 6, "6/67/Forrest_Gump_poster.jpg",
     "Un om simplu trece, aproape din întâmplare, prin marile momente ale Americii secolului XX."),
    ("F06", "Inception", "Începutul", 2010, "Christopher Nolan", "SF", 148, 160, 839, 8.8, 4, "2/2e/Inception_%282010%29_theatrical_poster.jpg",
     "Un hoț care fură secrete din vise primește misiunea inversă: să planteze o idee."),
    ("F07", "Fight Club", "Clubul luptătorilor", 1999, "David Fincher", "Dramă", 139, 63, 102, 8.8, 0, "f/fc/Fight_Club_poster.jpg",
     "Un funcționar cu insomnii și un vânzător de săpun pornesc un club de lupte clandestin."),
    ("F08", "The Matrix", "Matrix", 1999, "Lana și Lilly Wachowski", "SF", 136, 63, 473, 8.7, 4, "d/db/The_Matrix.png",
     "Un programator află că lumea în care trăiește este o simulare."),
    ("F09", "Interstellar", "Interstellar", 2014, "Christopher Nolan", "SF", 169, 165, 783, 8.7, 1, "b/bc/Interstellar_film_poster.jpg",
     "O echipă de astronauți caută prin altă galaxie o nouă casă pentru omenire."),
    ("F10", "Parasite", "Parazitul", 2019, "Bong Joon Ho", "Thriller", 132, 11.4, 258, 8.5, 4, "5/53/Parasite_%282019_film%29.png",
     "O familie săracă se infiltrează, pe rând, în casa unei familii bogate."),
    ("F11", "Titanic", "Titanic", 1997, "James Cameron", "Dramă", 195, 200, 2264, 7.9, 11, "1/18/Titanic_%281997_film%29_poster.png",
     "O poveste de dragoste pe vasul care s-a scufundat în 1912, la primul său drum."),
    ("F12", "Avatar", "Avatar", 2009, "James Cameron", "SF", 162, 237, 2924, 7.9, 3, "d/d6/Avatar_%282009_film%29_poster.jpg",
     "Un fost soldat ajunge pe luna Pandora și trece de partea localnicilor."),
    ("F13", "The Lord of the Rings: The Return of the King", "Stăpânul inelelor: Întoarcerea regelui", 2003, "Peter Jackson", "Fantasy", 201, 94, 1149, 9.0, 11, "4/48/Lord_Rings_Return_King.jpg",
     "Frodo și Sam ajung la Muntele Osândei, în timp ce oamenii apără ultimul regat."),
    ("F14", "Gladiator", "Gladiatorul", 2000, "Ridley Scott", "Acțiune", 155, 103, 466, 8.5, 5, "f/fb/Gladiator_%282000_film_poster%29.png",
     "Un general roman trădat devine gladiator și își caută dreptatea în arenă."),
    ("F15", "Spirited Away", "Călătoria lui Chihiro", 2001, "Hayao Miyazaki", "Animație", 124, 19.2, 396, 8.6, 1, "d/db/Spirited_Away_Japanese_poster.png",
     "O fetiță rămâne prinsă într-o lume a spiritelor și trebuie să-și salveze părinții."),
    ("F16", "The Lion King", "Regele Leu", 1994, "Roger Allers, Rob Minkoff", "Animație", 88, 45, 979, 8.5, 2, "3/3d/The_Lion_King_poster.jpg",
     "Puiul de leu Simba fuge de acasă și se întoarce să-și ia locul de rege."),
    ("F17", "Joker", "Joker", 2019, "Todd Phillips", "Dramă", 122, 55, 1079, 8.4, 2, "e/e1/Joker_%282019_film%29_poster.jpg",
     "Un comedian ratat din Gotham alunecă, pas cu pas, în violență."),
    ("F18", "Oppenheimer", "Oppenheimer", 2023, "Christopher Nolan", "Dramă", 180, 100, 983, 8.3, 7, "4/4a/Oppenheimer_%28film%29.jpg",
     "Povestea fizicianului care a condus construcția primei bombe atomice."),
    ("F19", "Barbie", "Barbie", 2023, "Greta Gerwig", "Comedie", 114, 145, 1448, 6.8, 1, "0/0b/Barbie_2023_poster.jpg",
     "Barbie pleacă din lumea perfectă a păpușilor în lumea reală."),
    ("F20", "Dune: Part Two", "Dune: Partea a doua", 2024, "Denis Villeneuve", "SF", 165, 190, 715, 8.5, 2, "5/52/Dune_Part_Two_poster.jpeg",
     "Paul Atreides se alătură fremenilor de pe Arrakis și pornește războiul împotriva Harkonnenilor."),
    ("F21", "Everything Everywhere All at Once", "Totul, peste tot, deodată", 2022, "Daniel Kwan, Daniel Scheinert", "SF", 139, 14.3, 143, 7.8, 7, "1/1e/Everything_Everywhere_All_at_Once.jpg",
     "Proprietara unei spălătorii descoperă că poate sări între universuri paralele."),
    ("F22", "Avengers: Endgame", "Răzbunătorii: Sfârșitul jocului", 2019, "Anthony și Joe Russo", "Acțiune", 181, 356, 2799, 8.4, 0, "0/0d/Avengers_Endgame_poster.jpg",
     "Supereroii rămași încearcă să anuleze dispariția a jumătate din univers."),
    ("F23", "Jurassic Park", "Parcul Jurasic", 1993, "Steven Spielberg", "Aventură", 127, 63, 1030, 8.2, 3, "e/e7/Jurassic_Park_poster.jpg",
     "Un parc cu dinozauri clonați scapă de sub control chiar înainte de deschidere."),
    ("F24", "Back to the Future", "Înapoi în viitor", 1985, "Robert Zemeckis", "SF", 116, 19, 393, 8.5, 1, "d/d2/Back_to_the_Future.jpg",
     "Un adolescent ajunge cu o mașină a timpului în 1955 și riscă să nu se mai nască."),
    ("F25", "Coco", "Coco", 2017, "Lee Unkrich", "Animație", 105, 175, 823, 8.4, 2, "9/98/Coco_%282017_film%29_poster.jpg",
     "Un băiat care visează la muzică ajunge în Țara Morților, de Ziua Morților."),
    ("F26", "4 Months, 3 Weeks and 2 Days", "4 luni, 3 săptămâni și 2 zile", 2007, "Cristian Mungiu", "Dramă", 113, 0.8, 9.8, 7.9, 0, "4/41/4months3weeks%262days_poster.jpg",
     "Două studente în România anului 1987, o zi și o noapte. Palme d'Or la Cannes."),
]

HEAD_FILME = ["ID_Film", "Titlu", "Titlu_RO", "An", "Regizor", "Gen", "Durata_Min", "Buget_Mil_USD",
              "Incasari_Mil_USD", "Nota_IMDb", "Oscaruri", "Poster_URL", "Descriere"]
HEAD_PROI = ["ID_Proiectie", "ID_Film", "Data", "Ora", "Sala", "Locuri", "Bilete_Vandute", "Pret_Bilet_RON"]

SALI = {"Sala Mare": (220, 35), "Sala Studio": (90, 28)}
ORE = {"Sala Mare": ["17:00", "20:30"], "Sala Studio": ["16:00", "18:30", "21:00"]}
START = date(2026, 11, 6)  # festivalul ține 10 zile


def filme_rows():
    return [[f[0], f[1], f[2], f[3], f[4], f[5], f[6], f[7], f[8], f[9], f[10], WIKI + f[11], f[12]] for f in FILME]


def proiectii():
    rnd = random.Random(1212)
    ids = [f[0] for f in FILME]
    nota = {f[0]: f[9] for f in FILME}
    slots = [(START + timedelta(days=d), s, o) for d in range(10) for s in SALI for o in ORE[s]]
    # fiecare film rulează de cel puțin 1 ori; sloturile rămase merg spre filmele cu nota mai mare
    order = ids[:]
    rnd.shuffle(order)
    while len(order) < len(slots):
        order.append(rnd.choices(ids, weights=[(nota[i] - 6) ** 2 for i in ids])[0])
    rnd.shuffle(order)
    rows = []
    for n, ((d, sala, ora), fid) in enumerate(zip(slots, order), start=1):
        loc, pret = SALI[sala]
        weekend = d.weekday() >= 4
        cerere = 0.42 + (nota[fid] - 6.8) * 0.16 + (0.16 if weekend else 0) + (0.08 if ora >= "20:00" else 0)
        bilete = max(8, min(loc, round(loc * (cerere + rnd.uniform(-0.12, 0.12)))))
        rows.append([f"P{n:03d}", fid, d, ora, sala, loc, bilete, pret])
    return rows


def write_xlsx(path, sheet, head, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append(r)
    for col, w in zip("ABCDEFGHIJKLM", [8, 34, 34, 7, 26, 11, 11, 14, 16, 10, 9, 70, 80]):
        ws.column_dimensions[col].width = w
    if sheet == "Fact_Proiectii":
        for c in ws["C"][1:]:
            c.number_format = "dd.mm.yyyy"
    wb.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    fr = filme_rows()
    pr = proiectii()
    write_xlsx(OUT / "Dim_Filme.xlsx", "Dim_Filme", HEAD_FILME, fr)
    write_xlsx(OUT / "Fact_Proiectii.xlsx", "Fact_Proiectii", HEAD_PROI, pr)
    with open(OUT / "Dim_Filme.csv", "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh, delimiter=";")
        w.writerow(HEAD_FILME)
        w.writerows(fr)
    js = {
        "filme": [dict(zip(HEAD_FILME, r)) for r in fr],
        "proiectii": [[r[1], r[2].isoformat(), r[4], r[5], r[6], r[7]] for r in pr],
    }
    (ROOT / "js" / "s12-date.js").write_text(
        "/* Generat de assets/s12/generate_filme_s12.py — nu editați de mână.\n"
        "   proiectii = [ID_Film, Data, Sala, Locuri, Bilete_Vandute, Pret_Bilet_RON]. */\n"
        "window.S12_DATE = " + json.dumps(js, ensure_ascii=False) + ";\n",
        encoding="utf-8",
    )
    print(len(fr), "filme,", len(pr), "proiecții,", sum(r[6] for r in pr), "bilete")


if __name__ == "__main__":
    main()
